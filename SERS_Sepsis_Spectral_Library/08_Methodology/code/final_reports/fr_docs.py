"""Builds the nine final Word reports. Run from this folder: python fr_docs.py"""
import os,glob,shutil,numpy as np,pandas as pd
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import fr_data as D,fr_figs as G
OUT=r"C:\Users\sukes\Downloads\sukesh\Final_Submission"
DIRS={'bulk':'01_Bulk_Experimental_Reports','sim':'02_Molecular_Simulation_Reports','cmp':'03_Simulation_vs_Bulk_Reports'}
MOLS=('NAM','LTA','LPS')
LONG={'NAM':'N-acetylmuramic acid (NAM)','LTA':'lipoteichoic acid (LTA)','LPS':'lipopolysaccharide (LPS)'}
WINS={'NAM':[(430,600),(780,1000),(1000,1280)],'LPS':[(380,520),(720,900),(960,1130)],'LTA':[(380,600),(715,900),(980,1120),(1430,1490),(1690,1790)]}
FONT='Arial'
# ---------------------------------------------------------------- document helpers
class Doc:
    def __init__(s,title,subtitle):
        s.d=Document(); sec=s.d.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
        for a in ('left_margin','right_margin'): setattr(sec,a,Cm(2.2))
        sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2.0)
        st=s.d.styles['Normal']; st.font.name=FONT; st.font.size=Pt(10.5); st.element.rPr.rFonts.set(qn('w:eastAsia'),FONT)
        st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.12
        for n,sz in (('Heading 1',14),('Heading 2',11.5)):
            h=s.d.styles[n]; h.font.name=FONT; h.font.size=Pt(sz); h.font.bold=True; h.font.color.rgb=RGBColor(0,0,0); h.element.rPr.rFonts.set(qn('w:eastAsia'),FONT)
            h.paragraph_format.space_before=Pt(14 if n=='Heading 1' else 9); h.paragraph_format.space_after=Pt(5); h.paragraph_format.keep_with_next=True
        s.fig=0; s.tab=0; s.n1=0; s.n2=0
        p=s.d.add_paragraph(); r=p.add_run(title); r.bold=True; r.font.size=Pt(19); p.paragraph_format.space_after=Pt(2)
        p=s.d.add_paragraph(); r=p.add_run(subtitle); r.font.size=Pt(11.5); r.font.color.rgb=RGBColor(0x40,0x40,0x40); p.paragraph_format.space_after=Pt(10)
        s._rule(); s._footer()
    def _rule(s):
        p=s.d.add_paragraph(); pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); bt=OxmlElement('w:bottom'); bt.set(qn('w:val'),'single'); bt.set(qn('w:sz'),'8'); bt.set(qn('w:color'),'000000'); b.append(bt); pPr.append(b); p.paragraph_format.space_after=Pt(4)
    def _footer(s):
        p=s.d.sections[0].footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for t,txt in (('begin',None),('instr',' PAGE '),('end',None)):
            r=p.add_run(); r.font.size=Pt(9)
            if t=='instr': e=OxmlElement('w:instrText'); e.set(qn('xml:space'),'preserve'); e.text=txt
            else: e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),t)
            r._r.append(e)
    def h1(s,t): s.n1+=1; s.n2=0; s.d.add_heading(f'{s.n1}. {t}',1)
    def h2(s,t): s.n2+=1; s.d.add_heading(f'{s.n1}.{s.n2} {t}',2)
    def p(s,t,bold_lead=None):
        p=s.d.add_paragraph()
        if bold_lead: r=p.add_run(bold_lead+' '); r.bold=True
        p.add_run(t); return p
    def b(s,t,lead=None):
        p=s.d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3)
        if lead: r=p.add_run(lead+' '); r.bold=True
        p.add_run(t)
    def figure(s,path,cap,w=16.4):
        s.fig+=1; s.d.add_picture(path,width=Cm(w)); pp=s.d.paragraphs[-1]; pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; pp.paragraph_format.keep_with_next=True; pp.paragraph_format.space_after=Pt(2)
        c=s.d.add_paragraph(); r=c.add_run(f'Figure {s.fig}. '); r.bold=True; r.font.size=Pt(9.5); r2=c.add_run(cap); r2.font.size=Pt(9.5); c.paragraph_format.space_after=Pt(10)
    def table(s,head,rows,widths,cap=None,fs=9):
        if cap:
            s.tab+=1; c=s.d.add_paragraph(); r=c.add_run(f'Table {s.tab}. '); r.bold=True; r.font.size=Pt(9.5); r2=c.add_run(cap); r2.font.size=Pt(9.5); c.paragraph_format.keep_with_next=True; c.paragraph_format.space_after=Pt(3)
        t=s.d.add_table(rows=1,cols=len(head)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for i,x in enumerate(head):
            c=t.rows[0].cells[i]; c.text=''; r=c.paragraphs[0].add_run(x); r.bold=True; r.font.size=Pt(fs)
            sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'D9D9D9'); c._tc.get_or_add_tcPr().append(sh)
        trPr=t.rows[0]._tr.get_or_add_trPr(); th=OxmlElement('w:tblHeader'); th.set(qn('w:val'),'true'); trPr.append(th)
        for row in rows:
            cs=t.add_row().cells
            for i,x in enumerate(row):
                cs[i].text=''; para=cs[i].paragraphs[0]; para.paragraph_format.space_after=Pt(1); para.add_run(str(x)).font.size=Pt(fs)
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        s.d.add_paragraph().paragraph_format.space_after=Pt(4)
    def save(s,path): s.d.save(path)
# ---------------------------------------------------------------- facts
def marker_table(m):
    f={'NAM':'03_Reference_Spectra/NAM/NAM_Reference_Marker_Peaks.csv','LPS':'03_Reference_Spectra/LPS/LPS_Reference_Marker_Peaks.csv','LTA':'03_Reference_Spectra/LTA/LTA_Reference_20261028_Pooled_Marker_Peaks.csv'}[m]
    t=pd.read_csv(D.LIB+'/'+f,comment='#'); out={}
    for _,r in t.iterrows():
        if m=='NAM': pos,sd,snr,det,role=r.pos_mean,r.pos_sd,r.snr_mean,f'{int(r.n_detected)}/{int(r.n)}',r.role; gz=bool(r.glass_zone)
        elif m=='LPS': pos,sd,snr,det,role=r.pos_mean,r.pos_sd,r.snr_median,r.detected,r['class']; gz=False
        else: pos,sd,snr,det,role=r.pos_mean,r.pos_sd,r.snr_median,r.detected,r['class']; gz=bool(r.glass_zone)
        out[round(float(pos),1)]=dict(sd=float(sd),snr=float(snr),det=str(det),role=str(role),gz=gz)
    return out
def facts(m):
    mem=pd.read_csv(D.LIB+'/'+D.MEMBERS[m]); k=D.peaks(m); mt=marker_table(m)
    cnt={'NAM':311,'LPS':51,'LTA':99}[m]
    return dict(n=len(mem),peaks=k,mt=mt,cnt=cnt)
def mk(mt,pos):
    best=min(mt,key=lambda a:abs(a-pos)); return mt[best]
def cross_r():
    out={}
    refs={m:D.exp_ref(m)[0] for m in MOLS}; G_=~((D.X>=D.GLASS[0])&(D.X<=D.GLASS[1]))
    for a,b in (('NAM','LPS'),('NAM','LTA'),('LPS','LTA')): out[(a,b)]=float(np.corrcoef(refs[a][G_],refs[b][G_])[0,1])
    return out
CR=cross_r()
SETTING={'NAM':'90 % of maximum laser power, 30 s integration, 5 accumulations','LPS':'60 % of maximum laser power, 40 s integration, 5 accumulations','LTA':'80 % of maximum laser power, 30 s with 3 accumulations or 45 s with 2 accumulations'}
SAMPLE={'NAM':'N-acetylmuramic acid powder (Sigma-Aldrich A3007; a mixture of the alpha and beta forms)','LPS':'lipopolysaccharide from Escherichia coli O111:B4 (Sigma-Aldrich L2630)','LTA':'lipoteichoic acid from Staphylococcus aureus (Sigma-Aldrich L2515)'}
ROLE={'NAM':'a sugar unit of the bacterial cell-wall peptidoglycan (a Gram-positive and Gram-negative cell-wall marker)','LPS':'the main lipid-sugar molecule of the outer membrane of Gram-negative bacteria (endotoxin), a trigger of sepsis','LTA':'a phosphate-and-glycerol polymer anchored in the membrane of Gram-positive bacteria, the Gram-positive counterpart of LPS'}
def intro_plain(doc):
    doc.p("A biomarker is a substance that can indicate a biological condition or process. A Raman spectrum is a pattern of signals (peaks) that depends on the chemical bonds in a molecule, so it can act as a fingerprint. A reference spectrum is a pattern of known origin, measured or calculated, against which other spectra are compared.")
# ---------------------------------------------------------------- BULK REPORT
def bulk(m):
    F=facts(m); mem=D.exp_members(m); res=D.match(m)
    ptab=D.peaks(m); mt=F['mt']
    figs={'full':G.bulk_full(m)[0],'proc':G.bulk_processing(m),'zoom':G.bulk_zoom(m,WINS[m][:3]),'glass':G.bulk_glass(m),'sel':G.bulk_select(m)}
    snr_med=np.median([mk(mt,p)['snr'] for p,_ in ptab]); main=[p for p,r in ptab if r=='main']
    d=Doc(f'Experimental Bulk Raman Spectra of {LONG[m][0].upper()+LONG[m][1:]}','Reference spectrum of the bulk pure biomarker')
    d.h1('Executive summary')
    n=F['n']
    if m=='NAM':
        d.p(f"N-acetylmuramic acid (NAM) is {ROLE[m]}. Its Raman fingerprint was established from {n} spectra of the bulk pure biomarker recorded at one setting ({SETTING[m]}), out of {F['cnt']} NAM spectra measured at many powers and integration times. The spectrum is dominated by a group of sharp peaks between 450 and 1130 cm⁻¹ at {', '.join(f'{p:.0f}' for p in main)} cm⁻¹ (main peaks) and is highly reproducible: every main peak is present in all {n} spectra and the signal-to-noise ratio of the strongest peaks is above 40.")
    elif m=='LPS':
        d.p(f"Lipopolysaccharide (LPS) is {ROLE[m]}. Its Raman fingerprint was established from {n} spectra of the bulk pure biomarker recorded at one setting ({SETTING[m]}). The spectrum has one clearly strongest peak at 425 cm⁻¹ and weaker peaks near 766, 861, 996 and 1080 cm⁻¹. The 425 cm⁻¹ peak is present in every spectrum; the weaker peaks are present in 8 to 10 of the {n}.")
    else:
        d.p(f"Lipoteichoic acid (LTA) is {ROLE[m]}. All {F['cnt']} available LTA spectra, recorded at nine laser-power and integration settings, were examined. A reference spectrum was built from the {n} spectra of highest quality, all at 80 % of the maximum laser power. The strongest LTA peak is at 1092 cm⁻¹, with further main peaks at 539 and 862 cm⁻¹ and supporting peaks at 418, 723, 799, 1010, 1460 and 1751 cm⁻¹.")
    d.p(f"A strong band near 1367 cm⁻¹ and its neighbours (1308 to 1455 cm⁻¹) come from the glass slide, not from the molecule; this zone is shaded in all figures and is not used as evidence. The spectrum is distinct from the other two molecules (correlation with {'LTA' if m=='NAM' else 'NAM'}: {CR[tuple(sorted((m,'LTA' if m=='NAM' else 'NAM'),key=['NAM','LPS','LTA'].index))]:.2f}; with {'LPS' if m!='LPS' else 'LTA'}: {CR[tuple(sorted((m,'LPS' if m!='LPS' else 'LTA'),key=['NAM','LPS','LTA'].index))]:.2f}; 1.0 would mean identical).")
    d.h1('Objective')
    d.p(f"To record a reliable, repeatable Raman spectrum of {LONG[m]} as a bulk pure biomarker, to identify its characteristic peaks, and to confirm that those peaks come from the molecule and not from the glass background of the measurement. This spectrum is the experimental reference for comparison with the molecular simulation and for later detection work.")
    d.h1('Data and approach')
    d.h2('What was measured')
    d.p(f"Sample: {SAMPLE[m]}. The bulk pure biomarker was measured with a 785 nm Raman instrument (B&W Tek i-Raman Plus). Laser power is given as a percentage of the instrument maximum. Reference condition: {SETTING[m]}. Each file is one spectrum, averaged over its accumulations; the number of independent measurement spots is stated where it matters.")
    d.h2('How the spectra were processed')
    d.p("Only steps that were actually applied are listed. Each is intended to make real peaks easier to see without changing their positions.")
    d.b("The instrument background (dark signal) was subtracted, and two defective detector pixels were repaired by interpolation.","Background subtraction:")
    d.b("Narrow random spikes (for example cosmic rays) were removed.","Spike removal:")
    d.b("A Savitzky–Golay filter (window of 11 points) reduced random fluctuation while keeping peak shapes.","Smoothing:")
    d.b("The slowly changing fluorescence-type background was estimated (asymmetric least-squares method) and subtracted, leaving only the peaks.","Baseline correction:")
    d.b("Each spectrum was scaled to run from 0 to 1 so spectra of different overall brightness can be compared by shape. Normalised plots show shape only; heights in counts are used where stated.","Normalisation:")
    d.p("The range 400 to 1800 cm⁻¹ was used. Correlations between spectra exclude the glass zone (1308 to 1455 cm⁻¹).")
    d.h1('Results')
    d.h2('The reference spectrum')
    d.figure(figs['full'],f"Raman spectra of bulk pure {LONG[m]}. Thin lines: each of the {n} spectra; bold line: mean; shaded band: ± 1 standard deviation. Numbers mark the main and supporting peaks. Grey: glass-slide zone.")
    d.p(f"The individual spectra overlap closely, which shows that the pattern is a property of the material and not of one measurement spot. "+{'NAM':"The peaks are narrow, which is typical of a crystalline sugar derivative.",'LPS':"The peaks are broader and weaker than for NAM, as expected for a large, heterogeneous molecule; spot-to-spot variation in height is larger, so intensities are compared as ratios to the 425 cm⁻¹ peak.",'LTA':"Peaks are broad, as expected for a polymer; heights vary by about 25 to 35 % between spots, but peak positions are stable (standard deviation about 1 to 2 cm⁻¹)."}[m])
    d.h2('Effect of processing')
    d.figure(figs['proc'],"One measured spectrum before and after processing. Top: the measured signal (grey), the signal after spike removal and smoothing (colour) and the estimated background (dashed). Bottom: the final spectrum after the background is removed.")
    d.p("Smoothing lowers the random fluctuation, and baseline removal flattens the rising background so that peak heights can be read from zero. Peak positions do not change.")
    d.h2('Main peaks')
    d.figure(figs['zoom'],"Close-ups of the main peaks. Thin lines: individual spectra; bold line: mean; band: ± 1 standard deviation. Numbers give the mean peak position.")
    rows=[]
    for p,role in sorted(ptab):
        if p<1308 or p>1455:
            info=mk(mt,p); key=min(D.EXPECT[m],key=lambda a:abs(a-p)); desc=D.EXPECT[m][key][1]
            rows.append([f"{p:.0f} ± {info['sd']:.1f}",'main' if role=='main' else 'supporting',desc,f"present in {info['det']} spectra; signal-to-noise {info['snr']:.0f}"])
    d.table(['Raman shift (cm⁻¹)','Peak role','Tentative assignment','Evidence from the measurements'],rows,[3.0,2.3,6.1,5.2],f"Peaks of the {m} reference spectrum (outside the glass zone). Assignments are tentative and based on general group frequencies of the molecular groups present.")
    d.h2('Comparison with the glass slide')
    d.figure(figs['glass'],f"The bulk sample compared with the glass blank. "+({'NAM':'Both curves are normalised; the glass slide was measured at low laser power because at the power used for the sample the glass band exceeds the detector range. ','LPS':'Top: signal in counts, same settings. Bottom: both normalised. ','LTA':'Top: signal in counts, same settings. Bottom: both normalised. '}[m])+f"Grey: the glass zone; dotted lines mark the {m} peaks.")
    gl={'NAM':"The glass slide contributes a strong broad band at 1308 to 1455 cm⁻¹ and weak, broad features elsewhere. None of the NAM main peaks coincides with a glass feature, so they belong to NAM.",'LPS':"The glass slide gives its strong band at 1308 to 1455 cm⁻¹ and little else. The 425 cm⁻¹ peak is absent from the blank slides and is therefore from LPS; the 1080 cm⁻¹ peak overlaps weakly with the glass background (about 14 % of its area), so it is treated with more caution.",'LTA':"The glass slide gives its strong band at 1308 to 1455 cm⁻¹. The LTA main peaks (539, 862 and 1092 cm⁻¹) are not present in the blank slides; a peak that appears near 1328 cm⁻¹ in the LTA spectra falls inside the glass zone and is not used."}[m]
    d.p(gl)
    d.h2('Choice of measurement condition')
    sel={'NAM':"Signal-to-noise, background ripple and agreement with the reference mean were compared for spectra recorded at 60, 70 and 80 % laser power, and the reference was recorded at 90 %. Higher power gives stronger peaks; the reference condition (90 %, 30 s, 5 accumulations) was kept because it gave the strongest peaks (signal-to-noise of the 929 cm⁻¹ peak about 50) without saturating the detector in the region of interest, and the lower-power spectra reproduce the same peak pattern (see figure).",
         'LPS':"Three conditions were compared (figure). Raising the integration time from 30 s to 40 s at 60 % laser power raised the signal-to-noise of the 425 cm⁻¹ peak (median about 21 against 15) with the same low ripple (about 16 to 18 counts). Raising the power to 70 % did not improve the peak but increased the ripple pattern about six-fold (about 115 counts), and at 80 % the ripple was larger still, so 60 % with 40 s and 5 accumulations was chosen. Peak heights still vary three- to five-fold between spots, which is a property of this heterogeneous material.",
         'LTA':"All 99 LTA spectra were compared on four criteria: signal-to-noise of the 1092 cm⁻¹ peak, background ripple, similarity to the pooled mean, and whether the background (dark) signal was free of sample light. At 80 % laser power the main peak is strong (median signal-to-noise about 28 to 29), the ripple is low (about 47 to 50 counts) and agreement with the mean is high. Two integration settings at 80 % (45 s × 2 and 30 s × 3 accumulations, equal total exposure) give the same pattern (correlation 0.98) and were pooled. The 30 s × 3 setting is preferred for new measurements because all seven spectra recorded with a fresh background reading passed every check. Settings at 60 % or below have weaker peaks, and the 70 % settings give similar or lower signal-to-noise with either higher ripple or poorer agreement with the mean."}[m]
    d.figure(figs['sel'],"Signal-to-noise, background ripple and similarity to the reference for the settings compared. Each point is one spectrum"+(": filled points have a reliable background reading; open points have a background that contained sample light." if m=='LTA' else "."))
    d.p(sel)
    d.h1('Interpretation')
    interp={'NAM':f"The NAM fingerprint is the clearest of the three: seven main peaks (541, 829, 871, 929, 955, 1023 and 1127 cm⁻¹) appear in every spectrum with high signal-to-noise and a position scatter of under 1 cm⁻¹. The peaks lie where sugar-ring skeletal, C–O–C and C–C stretching modes are expected. Because no pure-compound Raman reference for NAM was found in the literature, assignments rest on group frequencies and on the calculation (separate reports) and remain tentative. The weaker peaks at 1249 and 1589 cm⁻¹ are reproducible but have no firm assignment.",
            'LPS':"LPS gives a recognisable but weaker and more variable spectrum. The 425 cm⁻¹ peak is the strongest and the most reproducible, but it has no assignment in the published LPS literature; it lies in the low-frequency skeletal region of carbohydrates. The 861, 996 and 1080 cm⁻¹ peaks are consistent with sugar C–O–C and C–O stretching and, for 1080 cm⁻¹, possibly the symmetric PO₂⁻ stretch. The spectrum is distinct from NAM and LTA, but because the material is a mixture of molecules of different size, the pattern is less sharp than for NAM.",
            'LTA':"The strongest LTA peak, 1092 cm⁻¹, falls where the symmetric stretch of the phosphodiester group (PO₂⁻) of the glycerol-phosphate backbone is expected, and the 799 cm⁻¹ peak is where O–P–O stretching is expected; both are consistent with, but do not prove, an LTA backbone. The 1460 cm⁻¹ peak is CH₂ bending of the lipid anchor and is not specific. The 862 and 539 cm⁻¹ peaks lie in the glycerol and sugar regions. The spectrum differs from NAM (correlation "+f"{CR[('NAM','LTA')]:.2f}"+") and moderately resembles LPS ("+f"{CR[('LPS','LTA')]:.2f}"+"), as both are phosphorylated glycolipids."}[m]
    d.p(interp)
    d.p("A strong signal alone does not identify a molecule. Here the identification rests on the combination of reproducible peak positions, the agreement of the whole pattern across spectra, absence of the peaks from the glass slide, and separation from the other two molecules.")
    d.h1('Conclusion and next step')
    nxt={'NAM':"The NAM reference spectrum is final. More spots at 90 % power can be added, and the 1249 and 1589 cm⁻¹ peaks would benefit from a firm assignment, which the molecular simulation reports address.",
         'LPS':"The LPS reference is established at 60 % power. The next useful step is to repeat it with a fresh background reading on a second sample vial, and to assign the 425 cm⁻¹ peak by comparison with a calculation of the lipid A core (see the simulation reports).",
         'LTA':"The LTA reference is established and agrees with the earlier preliminary measurements (correlation 0.95). The next useful steps are to confirm the reference with new spots at 80 % power, 30 s and 3 accumulations using a freshly recorded background, and to obtain the supplier's certificate of analysis for the lot, because its purity is not stated."}[m]
    d.p(nxt)
    path=os.path.join(OUT,DIRS['bulk'],f'{m}_Bulk_Spectral_Analysis.docx'); d.save(path); return path
# ---------------------------------------------------------------- SIMULATION REPORT
SIMINFO={
 'NAM':dict(model="the whole N-acetylmuramic acid molecule in its beta form (39 atoms: C₁₁H₁₉NO₈)",
   approach="density functional theory (B3LYP with dispersion correction) with a large basis set (6-311++G(d,p)), for the isolated molecule",
   stages=[["1","Early runs at smaller basis sets (6-311+G(d,p) and 6-31+G(d,p))","Showed the workflow worked; not used for comparison"],["2","Final calculation at the large basis set with tight geometry optimisation","Used for all comparisons: 111 vibrations, none unstable, 67 between 400 and 1800 cm⁻¹"],["3","Alpha form of the molecule calculated at the same level","Tested as a possible explanation for unmatched peaks and did not explain them"]],
   why="NAM is small enough (39 atoms) for a high-level calculation of the whole molecule. A larger basis set and a dispersion correction describe the hydrogen-bonded sugar ring more reliably, and the optimised structure is a true energy minimum, so every calculated vibration is real.",
   scaling="Calculated frequencies from this kind of calculation are slightly too high, so they were multiplied by 0.98 below 1800 cm⁻¹ and 0.967 above. Each vibration was drawn as a Lorentzian peak 10 cm⁻¹ wide.",
   limits=["The calculation is for one isolated molecule in its beta form; the sample is a mixture of alpha and beta forms and is a crystalline powder, where molecules hydrogen-bond to each other.","Calculated intensities (how strong each peak should be) are less reliable than calculated positions.","Only the beta form is used in the comparison; the alpha form lies about 3.6 kcal/mol higher in electronic energy for the isolated molecule."],
   nxt="Calculate the alpha and beta forms together and combine them in proportion to their expected populations, and model a small cluster of NAM molecules to capture hydrogen bonding in the powder."),
 'LPS':dict(model="the core of lipopolysaccharide: lipid A with two Kdo sugars, a whole-molecule model of 201 atoms (C₆₀H₉₈N₂O₃₉P₂, charge −4) in which the six long fatty-acid chains were shortened to 3–4 carbons",
   approach="a quick two-step calculation: a fast semi-empirical optimisation (PM6) followed by density functional theory (B3LYP) with a small basis set (3-21G), for the isolated molecule",
   stages=[["1","Five small fragments of LPS calculated separately (Kdo, glucosamine, heptose, myristic acid and phosphoric acid) and their spectra summed","Earlier approach; the sum ignores the bonds that join the pieces, so it was set aside"],["2","One whole-molecule model of lipid A with shortened chains","Final model used for all comparisons: 597 vibrations, none unstable, 354 between 400 and 1800 cm⁻¹"],["3","Higher-level calculation of the whole model (larger basis set with water)","Planned; not yet available"]],
   why="Real LPS carries long polysaccharide chains and is far too large to calculate. Lipid A with its Kdo sugars is the conserved core that contains the phosphate groups, the glucosamine sugars, the amide links and the glycosidic bond. Keeping all of these bonds in one molecule is more realistic than summing separate fragments. The fatty-acid chains were shortened so that the molecule could be calculated on the available computer.",
   scaling="Calculated frequencies were multiplied by 0.965, a typical value for this method; it was not fitted for the small basis set. Raman intensities were calculated for 785 nm light at room temperature, and each vibration was drawn as a Lorentzian peak 12 cm⁻¹ wide.",
   limits=["The basis set is small (3-21G), so positions are approximate and intensities distorted, especially for the charged phosphate groups.","The model lacks the long fatty-acid chains, the outer core sugars, the O-antigen chain and counter-ions, and was calculated without water.","The strong calculated peak near 1516 cm⁻¹ involves the amide groups and is probably overstated by the small basis set; it is not seen in the measurement."],
   nxt="Repeat the calculation with a larger basis set and water, add sodium or calcium counter-ions and longer chains, and analyse the motions behind the calculated peaks near 422 and 862 cm⁻¹, to decide what the unassigned 425 cm⁻¹ measured peak is."),
 'LTA':dict(model="a whole-molecule model of 93 atoms (C₂₃H₄₅N₃O₂₀P₂) with three glycerol-phosphate units joined by two phosphodiester links, a D-alanine ester on each outer glycerol and one alpha-N-acetylglucosamine on the central glycerol (net charge zero)",
   approach="a quick two-step calculation: a fast semi-empirical optimisation (PM6) followed by density functional theory (B3LYP) with a small basis set (3-21G), for the isolated molecule",
   stages=[["1","Three small fragments calculated separately (glycerol-phosphate, D-alanine and N-acetylglucosamine) and their spectra summed","Earlier approach; the sum ignores the bonds that join the pieces, so it was set aside"],["2","One whole-molecule model with the three substituent types","Final model used for all comparisons: 273 vibrations, none unstable, 165 between 400 and 1800 cm⁻¹"],["3","Higher-level calculation (larger basis set with water)","Started; not yet available"]],
   why="LTA is a polymer of about 23 glycerol-phosphate units on a lipid anchor, with D-alanine or sugar groups on the glycerols; no single structure exists for it. A short chain that carries all three kinds of group keeps the real bonds (phosphodiester, ester and the sugar link) in a calculable molecule.",
   scaling="Calculated frequencies were multiplied by 0.965, a typical value for this method; it was not fitted for the small basis set. Raman intensities were calculated for 785 nm light at room temperature, and each vibration was drawn as a Lorentzian peak 12 cm⁻¹ wide.",
   limits=["The basis set is small (3-21G), so positions are approximate and intensities distorted, especially for the charged phosphate and ammonium groups; calculated peaks above about 2400 cm⁻¹ (outside the range used here) are probably artefacts of internal hydrogen bonds.","Only three repeating units and no lipid anchor: end-of-chain effects are exaggerated and repeat-unit bands are under-represented.","No counter-ions or water; the geometry of the two glycerol arms was assumed rather than checked against the natural form."],
   nxt="Finish the higher-level calculation with water, then extend the chain to four or five units with counter-ions and the lipid anchor, and check which atoms move in the calculated peaks near 1085, 860 and 800 cm⁻¹.")}
def top_modes(m,n=8):
    d=D.sim(m); k=np.where(d['k'])[0]; idx=sorted(k,key=lambda i:-d['I'][i])[:n]; pk=[p for p,_ in D.peaks(m)]; rows=[]
    for i in sorted(idx,key=lambda i:d['f'][i]):
        near=[p for p in pk if abs(p-d['f'][i])<=15]
        rows.append([f"{d['f'][i]:.0f}",d['label'][i],f"{d['I'][i]:.2f}",f"yes ({near[0]:.0f})" if near else "not seen as a distinct peak"])
    return rows,idx
def assign_items(m,res,maxn=10):
    items=[];used=set()
    order=sorted([r for r in res if r['idx'] is not None],key=lambda r:(r['role']!='main',r['exp']))[:maxn]
    for r in sorted(order,key=lambda r:r['exp']): items.append(dict(idx=r['idx'],text=f"{r['exp']:.0f} measured / {r['calc']:.0f} calculated\n{r['label'].split(' (')[0]}"))
    return items
def simrep(m):
    S=SIMINFO[m]; res=D.match(m); d_=D.sim(m); rows,idx=top_modes(m)
    figs={'struct':G.struct_plain(m),'spec':G.sim_spectrum(m)}
    it=assign_items(m,res); figs['assign']=G.struct_assign(m,it,'sim_assign')
    d=Doc(f"Molecular Simulation of {LONG[m][0].upper()+LONG[m][1:]}",'Calculated Raman spectrum and assignment of the main vibrations')
    nk=int(d_['k'].sum())
    d.h1('Executive summary')
    d.p(f"A computer calculation of the Raman spectrum of {S['model']} was carried out to predict which Raman peaks this structure should give and which part of the molecule each peak comes from. The calculation gives {nk} vibrations between 400 and 1800 cm⁻¹. Its main product is a list of calculated peak positions with the atoms that move in each, which lets the measured peaks be assigned to groups such as the phosphate, the sugar ring or the ester carbonyl. The calculation reproduces peak positions more reliably than peak heights.")
    d.h1('Objective')
    d.p(f"Raman peaks depend on how the atoms of a molecule vibrate. The experimental spectrum of {LONG[m]} shows several reproducible peaks but their origin is not documented in the literature for the pure compound. The simulation was run to predict the Raman spectrum from the molecular structure, so that measured peaks can be linked to specific vibrations and the identification of the molecule is supported by an independent line of evidence.")
    d.h1('Molecular model and approach')
    d.p(f"Model: {S['model']}. Method: {S['approach']}. In plain terms, the computer first finds the lowest-energy shape of the molecule, then calculates how each group of atoms vibrates and how strongly each vibration would scatter laser light. A calculated spectrum is therefore a prediction of the Raman signals expected from the structure.")
    d.table(['Step','What was calculated','Outcome'],S['stages'],[1.2,8.4,7.0],"Calculations carried out and the one used.")
    d.h2('Why the final calculation is suitable')
    d.p(S['why'])
    d.h2('From vibrations to a spectrum')
    d.p(S['scaling']+" The nature of each vibration (for example phosphate stretch, ring deformation or CH bending) was assigned by measuring which atoms move most in it; these labels are tentative.")
    figs_struct=figs['struct']
    d.h1('Results')
    d.h2('Molecular structure')
    d.figure(figs_struct,"Structure of the model that was calculated.",w=15.5 if m!='LPS' else 16.4)
    d.h2('Calculated Raman spectrum')
    d.figure(figs['spec'],"Calculated Raman spectrum. Orange curve: all vibrations broadened to peaks; grey lines: the individual calculated vibrations (heights are relative Raman intensities); numbers: the six strongest.")
    d.h2('Important calculated peaks and their assignment')
    d.table(['Calculated peak (cm⁻¹)','Type of vibration','Relative intensity','Counterpart in the measured spectrum'],rows,[3.2,6.6,2.6,4.2],"The eight strongest calculated vibrations between 400 and 1800 cm⁻¹ (intensity 1.0 = strongest).")
    d.p("The strongest calculated peaks are not always the strongest measured peaks; Raman intensities from small calculations are only a rough guide. The comparison report matches calculated peaks to measured ones by position and by type of vibration.")
    d.h2('Assignment of the peaks seen in the measurement')
    d.figure(figs['assign'],"Atoms that move most in the calculated vibrations that correspond to measured peaks. Each colour is one vibration (key at right); the highlighted atoms carry most of the motion. Hydrogen atoms are assigned to the atom they are bonded to.",w=16.4)
    prow=[[f"{r['calc']:.0f}",f"{r['exp']:.0f}",r['label'],('confident' if r['status']=='agrees' else 'tentative')] for r in sorted(res,key=lambda r:r['exp']) if r['idx'] is not None]
    d.table(['Calculated (cm⁻¹)','Matching measured (cm⁻¹)','Assignment (type of vibration)','Confidence'],prow,[3.2,3.6,7.4,2.4],"Calculated vibrations matched to measured peaks. 'Confident' means the calculated peak lies within 10 cm⁻¹ of the measured peak and has the expected type of motion; 'tentative' means larger offset (see the comparison report).")
    d.h1('Interpretation and limitations')
    d.p("The calculation supports the assignment of the measured peaks to the vibration types listed above and shows which groups of the molecule are responsible. It does not by itself prove the identity of the sample.")
    for l in S['limits']: d.b(l)
    d.h1('Conclusion and next step')
    d.p(f"The calculation gives a usable prediction of the positions and types of the Raman vibrations of {LONG[m]} between 400 and 1800 cm⁻¹ and is the simulation used in the comparison with the measured spectrum. Next step: {S['nxt']}")
    path=os.path.join(OUT,DIRS['sim'],f'{m}_Molecular_Simulation_Report.docx'); d.save(path); return path
# ---------------------------------------------------------------- COMPARISON REPORT
def shift_reason(r):
    dl=r['delta']
    if dl<0: return ("The calculated peak lies at higher wavenumber than the measured one. This is the usual direction for a harmonic calculation of an isolated molecule, because hydrogen bonding and packing in the powder soften vibrations that the calculation does not include, and the scale factor is only approximate.")
    return ("The calculated peak lies at lower wavenumber than the measured one. A small basis set and isolated charged or hydrogen-bonded groups (no counter-ions or neighbouring molecules) can lower the calculated frequency, and the powder environment can raise the measured one.")
def cmprep(m):
    res=D.match(m); S=SIMINFO[m]; r_all=D.pearson(m); fig={}
    fig['full']=G.cmp_full(m,res); fig['zoom']=G.cmp_zoom(m,res,WINS[m]); fig['glass']=G.cmp_glass(m,res); fig['shift']=G.cmp_shift(m,res)
    it=assign_items(m,res); fig['assign']=G.struct_assign(m,it,'cmp_assign')
    ag=[r for r in res if r['status']=='agrees']; sh=[r for r in res if r['status']=='shifted']; no=[r for r in res if r['status']=='none']
    nmain=sum(1 for r in res if r['role']=='main')
    d=Doc(f"Calculated and Measured Raman Spectra of {LONG[m][0].upper()+LONG[m][1:]}",'Comparison of the simulation with the experimental bulk spectrum')
    d.h1('Executive summary')
    d.p(f"The calculated Raman spectrum of the {m} model was compared with the measured spectrum of the bulk pure biomarker (mean of {facts(m)['n']} spectra). Of the {len(res)} measured main and supporting peaks, {len(ag)} have a calculated vibration of the expected type within 10 cm⁻¹, {len(sh)} have one of the same type that is shifted by 10 to 50 cm⁻¹, and {len(no)} have no calculated counterpart of that type. The two spectra agree in where the peaks fall and in the type of vibration behind them, but not in overall shape (correlation {r_all:.2f}, glass zone excluded), because calculated intensities are unreliable. The calculation therefore supports the assignment of the measured peaks but does not by itself identify the molecule.")
    d.h1('Objective')
    d.p("To test whether the peaks measured for the bulk pure biomarker correspond to vibrations predicted for the molecular structure, and to explain peaks that are shifted or not reproduced, so the calculation can be used as supporting evidence for the identification of the molecule.")
    d.h1('How the spectra were compared')
    d.b("No shift was applied to either spectrum. The calculated frequencies carry only the single scale factor described in the simulation report (a general correction for the method, not tuned to this measurement).","Alignment:")
    d.b("Each spectrum was scaled to its own maximum between 400 and 1800 cm⁻¹, so only shape and position are compared, not absolute signal. In the close-ups each curve is scaled to its tallest peak in that window.","Normalisation:")
    d.b("The measured peaks are the main and supporting peaks of the reference spectrum (mean of repeated spectra). For each, the calculated vibration of the expected type (for example phosphate stretch, ring deformation, CH bending) nearest in position within 50 cm⁻¹ and with appreciable intensity was selected. 'Agrees' means a difference of 10 cm⁻¹ or less. 'Shifted' means 10 to 50 cm⁻¹ with the same type of vibration, so the peak is judged to come from the same bond or group but at a displaced position. 'Not reproduced' means no vibration of that type within 50 cm⁻¹.","Matching:")
    d.b("The glass zone (1308 to 1455 cm⁻¹) is shaded and not used.","Glass:")
    d.h1('Results')
    d.h2('Whole spectrum')
    d.figure(fig['full'],"Measured spectrum (top, mean ± 1 standard deviation) and calculated spectrum (bottom), each normalised. Dashed lines join peaks that correspond: green, agree within 10 cm⁻¹; orange, shifted but of the same vibration type; grey, no counterpart of that type.")
    d.p(f"The calculated spectrum has more peaks than the measurement, because it lists every vibration while the measured peaks are broadened and overlap, and the calculated heights do not follow the measured heights. The overall correlation is therefore low ({r_all:.2f}) and is not used as evidence of identity. The meaningful comparison is peak by peak.")
    d.h2('Peak-by-peak comparison')
    ST_={'agrees':'Agrees','shifted':'Shifted, same vibration','none':'Not reproduced'}
    rows=[[f"{r['exp']:.0f}",(f"{r['calc']:.0f}" if r['calc'] is not None else '—'),(f"{r['delta']:+.0f}" if r['delta'] is not None else '—'),(r['label'] if r['label'] else f"nearest calculated peak {r['near_calc']:.0f}: {r['near_label']}"),ST_[r['status']]] for r in sorted(res,key=lambda r:r['exp'])]
    d.table(['Measured (cm⁻¹)','Calculated (cm⁻¹)','Difference','Calculated vibration (type)','Result'],rows,[2.6,2.8,2.0,6.6,2.8],"Measured peaks and the corresponding calculated vibrations. Difference = measured minus calculated.")
    d.figure(fig['zoom'],"Close-ups of the important regions. Upper curve: measurement; lower curve: calculation. Dashed lines mark the measured peak (above) and the corresponding calculated peak (below); arrows and numbers give the size of the shift where it exceeds 10 cm⁻¹.")
    d.h2('Shifted peaks: same vibration, different position')
    if sh:
        d.p("The following peaks do not coincide in position, but the calculated vibration that lies nearest is of the same type as the one expected for the measured peak, so they are treated as the same peak and the same bond or group, displaced:")
        for r in sorted(sh,key=lambda r:r['exp']):
            d.b(f"measured {r['exp']:.0f}, calculated {r['calc']:.0f} ({r['delta']:+.0f} cm⁻¹): {r['label']}. Expected assignment: {r['desc']}. "+shift_reason(r)+(" The offset is large, so the correspondence is tentative." if abs(r['delta'])>25 else ""))
    else: d.p("No peak is shifted by more than 10 cm⁻¹ while keeping the same vibration type.")
    d.figure(fig['shift'],"Position difference (measured minus calculated) for every matched peak. Green band: within 10 cm⁻¹.")
    if no:
        d.h2('Peaks not reproduced')
        for r in no: d.b(f"measured {r['exp']:.0f}: {r['desc']}. No calculated vibration of this type lies within 50 cm⁻¹; the nearest calculated peak ({r['near_calc']:.0f}) is a different type ({r['near_label']}). This peak is not explained by the present model.")
    d.h2('Molecule or glass?')
    d.figure(fig['glass'],"Measured bulk sample, glass blank and calculated spectrum. Grey: glass zone; dotted lines: measured molecular peaks. The calculated spectrum is for the molecule only and contains no glass signal.")
    d.p("Measured peaks that fall outside the glass zone and are absent from the blank slide, and that have a calculated molecular vibration at or near the same position, are attributed to the molecule. Peaks in the glass zone are not compared.")
    d.h2('Assignment on the molecular structure')
    d.figure(fig['assign'],"Atoms that move most in the calculated vibrations that correspond to measured peaks. Colours link each calculated vibration to its measured peak (key at right).",w=16.4)
    extra={'NAM':"Many NAM vibrations involve the whole ring at once, so the highlighted atoms are broad groups rather than single bonds.",'LPS':"Highlighted atoms are heavy atoms only; the vibrations of this large molecule are spread over several groups.",'LTA':"The phosphate group (pink) is responsible for the calculated peak that corresponds to the strong 1092 cm⁻¹ measured peak."}[m]
    d.p(extra)
    d.h1('Interpretation')
    imp={'NAM':"Seven of the eleven measured NAM peaks have a calculated vibration of the expected ring-deformation or C–O / C–C stretching type within 10 cm⁻¹ (457, 487, 871, 929, 1023, 1127 and 1249 cm⁻¹). The 541 and 955 cm⁻¹ peaks are matched with shifts of 15 and 24 cm⁻¹ and the 829 cm⁻¹ main peak with a shift of 36 cm⁻¹ to a strong calculated ring stretching mode at 793 cm⁻¹; this correspondence is the least certain and is the main unexplained feature of the NAM spectrum, and the 1589 cm⁻¹ peak has no counterpart. Because the sample is a mixture of alpha and beta forms and the calculation is for the beta form of a single molecule, some displacement of the sugar-ring modes is expected. The agreement is useful to assign the measured peaks to sugar-ring vibrations; it is not a stand-alone identification.",
         'LPS':"All five measured LPS peaks have a calculated vibration of the expected type within 3 cm⁻¹. This includes the 425 cm⁻¹ peak, which the calculation places on a low-frequency skeletal deformation of lipid A, and the 861, 996 and 1080 cm⁻¹ peaks on sugar C–O and C–C stretching. The agreement should be read with care: the model has 354 vibrations between 400 and 1800 cm⁻¹, so a calculated peak is almost always close to any measured peak, and a test with randomly displaced calculated peaks reached the same level of agreement about every second time (p = 0.57). The measured 1080 cm⁻¹ peak is matched by a sugar C–O stretch rather than by a phosphate vibration. Together, this supports the assignment of the LPS peaks to sugar and phosphate backbone motions, but does not establish the origin of the 425 cm⁻¹ peak.",
         'LTA':"All measured LTA peaks except 1751 cm⁻¹ have a calculated vibration of the expected type within 7 cm⁻¹, including the strongest measured peak, 1092 cm⁻¹, which matches a phosphate P–O stretching vibration at 1085 cm⁻¹, and the 799 and 862 cm⁻¹ peaks, which match C–O and C–C stretching vibrations of the glycerol-phosphate backbone and sugar. The 1751 cm⁻¹ peak corresponds to a calculated ester C=O stretch at 1726 cm⁻¹ (shifted by 25 cm⁻¹): a lower calculated C=O frequency is likely related to the isolated, charged model with a small basis set, in which the carbonyl can interact with the nearby ammonium group; the powder environment is different. The nearest calculated peak to 1751 cm⁻¹ in position (1746 cm⁻¹) is an N–H motion and was not used, because it is not a C=O vibration. The calculated vibrations that match are weak in intensity, and the strongest calculated peaks (690 and 1490 to 1510 cm⁻¹) are not seen in the measurement, so the overall shape does not agree. In a test with randomly displaced calculated peaks (which counts the nearest calculated peak whatever its type), agreement at 5 cm⁻¹ was reached in fewer than 1 in 1000 trials, but at 10 cm⁻¹ about 1 in 14 (p = 0.07). The agreement therefore supports the phosphate, glycerol and ester assignments but does not identify LTA on its own."}[m]
    d.p(imp)
    d.h1('Conclusion and next step')
    d.p(f"The calculation and the measurement agree on peak positions and vibration types for {len(ag)+len(sh)} of {len(res)} {m} peaks, with the shifts explained by the isolated-molecule, small-model and gas-phase nature of the calculation. They do not agree on intensities or overall shape. The calculation is used as supporting evidence for peak assignment ('consistent with'), not as proof of identity; identification rests on the reproducible measured fingerprint and its separation from the other two molecules. Next step: {S['nxt']}")
    path=os.path.join(OUT,DIRS['cmp'],f'{m}_Simulation_vs_Bulk_Comparison.docx'); d.save(path); return path
if __name__=='__main__':
    for k in DIRS.values(): os.makedirs(os.path.join(OUT,k),exist_ok=True)
    for m in MOLS: print(bulk(m)); print(simrep(m)); print(cmprep(m))
