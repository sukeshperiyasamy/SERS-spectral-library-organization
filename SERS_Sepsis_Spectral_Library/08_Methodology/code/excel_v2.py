import sys,os,pickle,json,numpy as np,pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
sys.path.insert(0,os.path.expanduser('~/work')); from lib2 import *; from lit_v2 import *
B,P,A,G,GZ=pickle.load(open(os.path.expanduser('~/work/v2.pkl'),'rb')); PK,ALL,REP,BULK,COND,res=pickle.load(open(os.path.expanduser('~/work/an2.pkl'),'rb'))
H=Font(name='Arial',bold=True,color='FFFFFF',size=10); HF=PatternFill('solid',fgColor='1F4E79'); F=Font(name='Arial',size=10); thin=Side(style='thin',color='BFBFBF')
def sheet(wb,name,df,title,notes=(),widths=None):
  ws=wb.create_sheet(name); ws.cell(1,1,title).font=Font(name='Arial',bold=True,size=12); r0=2
  for n in notes: ws.cell(r0,1,n).font=Font(name='Arial',italic=True,size=9,color='52514E'); r0+=1
  r0+=1
  for j,c in enumerate(df.columns,1):
    x=ws.cell(r0,j,c); x.font=H; x.fill=HF; x.alignment=Alignment(wrap_text=True,vertical='center')
  for i,row in enumerate(df.itertuples(index=False),r0+1):
    for j,v in enumerate(row,1):
      if isinstance(v,np.generic): v=v.item()
      if isinstance(v,bool): v='Yes' if v else 'No'
      if isinstance(v,float) and np.isnan(v): v=None
      c=ws.cell(i,j,v); c.font=F; c.border=Border(bottom=thin)
      s=str(v)
      if s.startswith('Excluded') or 'Contaminated' in s or s.startswith('Saturated'): c.font=Font(name='Arial',size=10,color='C00000')
      elif s.startswith(('Flagged','Signal not','Review','Abnormal')): c.font=Font(name='Arial',size=10,color='B26B00')
  ws.freeze_panes=ws.cell(r0+1,2); ws.auto_filter.ref=f"A{r0}:{get_column_letter(len(df.columns))}{r0+len(df)}"
  for j,c in enumerate(df.columns,1):
    w=(widths or {}).get(c) or min(55,max(9,len(str(c))*0.9+2,*(len(str(v))+2 for v in df[c].head(300))))
    ws.column_dimensions[get_column_letter(j)].width=w
  return ws
wb=Workbook(); ws=wb.active; ws.title='README'
L=[('SERS Sepsis Project — Bulk Powder Raman Spectral Library (NAM, LTA, LPS)',True),('Bulk (neat powder) samples measured directly on glass slides; B&W Tek i-Raman Plus BWS465-785H (BTC665N-785H-SYS), 785 nm; processed range 400–1800 cm-1.',False),
('Laser power is the instrument power level (%) from the file header; NOT converted to mW (sample power not measured). Objective: not recorded in the raw file headers (probe_name field empty).',False),
('Original raw files were never modified; 01_Raw_Data_Renamed contains byte-identical renamed copies (MD5-checked).',False),('',False),
('Naming: <Analyte>_Bulk_<YYYYMMDD>_P<power>pct_T<time>s_A<acc>_S<rep>.csv ; blanks: Glass_Blank_<YYYYMMDD>_P..pct_T..s_A.._S...csv ; replicate numbers are per date/condition in acquisition order.',False),
('Two NAM files were saved by the instrument in TXTR (semicolon) format; their renamed copies keep the .txtr extension.',False),('',False),
('Version 2.1 (23 Sep 2026): NAM 60 s re-audit, 70 % supporting group, final peak hierarchy (primary / supporting / glass-background / uncertain), LTA labelled PRELIMINARY (n = 4), LPS pending.',True),('Sheets: 1_File_Inventory | 2_QC | 3_Reference_Conditions | 4_Final_Marker_Table | 5_Excluded_Data | 6_Literature | 7_Literature_Bands | 8_All_Reproducible_Peaks | 9_Reproducibility | 10_Bulk_Intensity | 11_Future_Acquisition | 12_NAM_60s_ReAudit | 13_NAM_70pct_vs_Reference | 14_NAM_CrossDay_Detection',False),('',False),
('Key QC rules: Saturated (raw ≥ 65 000 counts) → excluded; SNR < 10 → Low SNR (kept, not used for reference); polystyrene-like pattern → contamination (excluded); correlation to class mean (glass zone 1308–1455 cm-1 excluded) < 0.85 → review.',False),
('SNR = height of strongest marker band (baseline-corrected) / SD of high-frequency noise (despiked minus SG-smoothed). Pearson correlation measures spectral reproducibility, not molecular identity.',False),
('Experimental peaks and literature peaks are reported in separate columns. Assignments are tentative unless stated; wording: "consistent with", not "proves".',False)]
for i,(t,b) in enumerate(L,1): ws.cell(i,1,t).font=Font(name='Arial',size=12 if i==1 else 10,bold=b)
ws.column_dimensions['A'].width=160
inv=A.merge(B[['orig','new_name_raw','cls','rep','qc_status','raw_folder']],on='orig',how='left',suffixes=('','_b'))
lab={'NAM':'NAM bulk','LTA':'LTA bulk','LPS':'LPS bulk','GLASS':'Glass blank','SUSPECT':'Unknown – labelled LTA, polystyrene-like','OUT_OF_SCOPE':'Not bulk – SERS/substrate experiment (out of scope)'}
inv['Biomarker / class']=inv.cls.map(lab); inv['Status']=inv.qc_status.fillna('Out of scope – not copied')
I1=inv[['orig','new_name_raw','Biomarker / class','date','power','int_s','acc','rep','spot_type','objective','laser_nm','rs_min','rs_max','npts','raw_max','Status','raw_folder']].copy()
I1.columns=['Original file (relative to sukesh folder)','New file name','Biomarker / class','Acquired','Power (%)','Integration (s)','Accumulations','Replicate','Spot','Objective','Laser (nm)','Raman shift min','Raman shift max','Data points','Max raw counts (full range)','Status','Library folder']
sheet(wb,'1_File_Inventory',I1.sort_values(['Biomarker / class','Acquired']),'Sheet 1 — File inventory (every raw file in the data folder)',['Out-of-scope files (SERS/R6G/AgNO3/HF/substrate studies) are listed for completeness but not copied into the library.'])
Q=B[['new_name_raw','orig','cls','snr','raw_max','raw_max_fp','ds_max_fp','ds_min_fp','baseline_level','noise','saturated','sat_fingerprint','nspikes','r_class_exglass','qc_status','qc_reason','in_reference']].copy()
Q.columns=['File name','Original file','Class','SNR','Max raw counts (full range)','Max raw counts (400–1800)','Max dark-subtracted (400–1800)','Min dark-subtracted (400–1800)','Baseline level (median, counts)','Noise SD (counts)','Saturated (any pixel)','Saturated in 400–1800','Spikes removed','r to class mean (glass zone excl.)','QC decision','Reason','Used in reference']
for c in ['SNR','Max dark-subtracted (400–1800)','Min dark-subtracted (400–1800)','Baseline level (median, counts)','Noise SD (counts)','r to class mean (glass zone excl.)']: Q[c]=Q[c].astype(float).round(2)
sheet(wb,'2_QC',Q,'Sheet 2 — Quality control (all bulk / blank / suspect spectra)',widths={'Reason':70})
C3=COND.copy(); C3.columns=['Biomarker','Power (%)','Integration (s)','Accumulations','n','n Valid','n saturated','Mean SNR','SNR SD','Mean strongest-marker intensity (counts)','Intensity CV (%)','Mean baseline (counts)','Mean max raw counts (400–1800)','Mean pairwise r','Dates','Reference condition']
for c in C3.columns[7:14]: C3[c]=C3[c].astype(float).round(2)
sheet(wb,'3_Reference_Conditions',C3.sort_values(['Biomarker','Mean SNR'],ascending=[True,False]),'Sheet 3 — Acquisition conditions and reference selection',
 ['Reference: NAM 90 % / 30 s / 5 acc (n = 15); LTA 80–95 % / 45 s / 2 acc (n = 4). LPS: none (pending glass-blank-controlled screening).','NAM 70 % / 60 s / 5 acc (n = 3): laser-line saturation only; fingerprint valid (class A). Retained as a SUPPORTING group, not merged into the reference. NAM 90 % / 60 s / 5 acc (n = 10): distorted (class C), excluded.'])
rows=[]
for _,p in PK.iterrows():
  cat,lp,diff,ass,src,tech,cit,conf=MLIT[(p.analyte,p.nominal)]
  bg='Inside glass zone (1308–1455 cm-1)' if p.glass_zone else ('Adjacent to glass zone; ~18 % of height glass-related (exploratory)' if (p.analyte=='LTA' and p.nominal==1459) else 'Outside glass zone')
  if cat=='Glass/background': bg='Glass-associated feature'
  rows.append([p.analyte,cat,f"{p.pos_mean:.1f}",f"{p.pos_mean:.1f} ± {p.pos_sd:.1f}",round(p.height_mean,0),round(p.height_sd,0),round(p.cv,1),round(p.snr_mean,1),round(p.prom_mean,0),f"{p.n_detected}/{p.n}",lp,diff,ass,cit,conf,bg,round(p.glass_corr_height,0)])
M4=pd.DataFrame(rows,columns=['Biomarker','Peak category','Experimental wavenumber (cm-1)','Mean ± SD (cm-1)','Mean intensity (counts, baseline-corr.)','Intensity SD','CV (%)','SNR','Prominence (counts)','Replicate frequency (detected / n)','Literature wavenumber / region (cm-1)','Difference (exp − lit)','Proposed assignment','Literature citation','Confidence','Background status','Height after exploratory glass-template correction'])
catorder={'Primary marker':0,'Supporting peak':1,'Glass/background':3}
M4['_o']=M4['Peak category'].map(lambda c:catorder.get(c,2)); M4=M4.sort_values(['Biomarker','_o']).drop(columns='_o')
sheet(wb,'4_Final_Marker_Table',M4,'Sheet 4 — FINAL marker table (bulk powder on glass slide; NAM reference n = 15; LTA preliminary reference n = 4)',
 ['Primary marker = reproducible in all reference replicates and across measurement days, SNR ≥ 20 (NAM) / ≥ 15 (LTA), outside or separable from the glass zone. Supporting = reproducible but weaker/less specific. Glass/background = glass-associated. Uncertain = observed, insufficient evidence.',
  'Literature: '+NOPURE+' Assignments are proposed from general Raman group-frequency tables (L5–L7) and are NOT compound-specific proof.'],widths={'Proposed assignment':50,'Background status':32})
E=B[B.qc_status!='Valid'][['new_name_raw','orig','cls','power','int_s','acc','qc_status','qc_reason','raw_folder']].copy()
E.columns=['File name','Original file','Class','Power (%)','Integration (s)','Acc','Decision','Reason','Library folder']
sheet(wb,'5_Excluded_Data',E.sort_values(['Decision','Class']),'Sheet 5 — Excluded, flagged and low-SNR spectra (none deleted)',widths={'Reason':80})
Lt=pd.DataFrame(REFS)[['id','cite','doi','biomarker','technique','source','verified']]; Lt.columns=['ID','Full citation','DOI','Biomarker','Technique','Organism / source','Access note']
sheet(wb,'6_Literature',Lt,'Sheet 6 — Literature used',['No peer-reviewed normal-Raman spectrum of pure N-acetylmuramic acid or isolated LTA powder was located in this search; NAM/LTA assignments rely on group frequencies (L5–L7).'],widths={'Full citation':90})
Lb=pd.DataFrame(LITBANDS,columns=['Biomarker / group','Ref ID','Technique','Reported band (cm-1)','Reported assignment'])
def cmp(r):
  exp={'LTA':[802,861,1092,1459],'Phosphodiester (LTA backbone)':[802,1092]}.get(r['Biomarker / group'],[])
  try: v=float(r['Reported band (cm-1)'])
  except: return 'region – see Sheet 4'
  if r['Biomarker / group']=='LPS': return 'No LPS experimental peak established' + (' | lies in glass zone' if 1308<=v<=1455 else '')
  if not exp: return ''
  d=min(exp,key=lambda e:abs(e-v)); return f'closest exp. {d} ({d-v:+.0f} cm-1)'+(' – consistent' if abs(d-v)<=12 else ' – no match')
Lb['Comparison with our data']=Lb.apply(cmp,axis=1)
Lb['Note']=Lb['Technique'].apply(lambda t:'SERS bands can differ from bulk Raman (surface selection rules)' if 'SERS' in t else ('Whole-cell context' if 'whole' in t else 'Group frequency'))
sheet(wb,'7_Literature_Bands',Lb,'Sheet 7 — Literature band list')
Al=ALL.copy(); Al.columns=['Biomarker','Peak (cm-1)','Replicates detected','n','Frequency (%)','Relative intensity','In glass zone']
sheet(wb,'8_All_Reproducible_Peaks',Al.round(2),'Sheet 8 — All peaks detected in reference replicates (prominence ≥ 5 × noise)',['Full list for traceability; the report highlights only the selected marker peaks.'])
R9=REP.round(3); R9.columns=['Biomarker','n','Mean pairwise r','Min pairwise r','Mean r to mean spectrum','Min r to mean spectrum']
sheet(wb,'9_Reproducibility',R9,'Sheet 9 — Spectral reproducibility (Pearson r of normalised spectra, 400–1800 cm-1)',['Correlation measures reproducibility, not molecular identity.'])
b=BULK.round(1); b.columns=['Biomarker','Condition','n','Max raw counts mean','Max raw counts SD','Max dark-subtracted mean','Max dark-subtracted SD','Baseline mean (counts)','Strongest marker intensity mean','Strongest marker intensity SD','CV (%)','SNR mean','SNR SD','Dates']
sheet(wb,'10_Bulk_Intensity',b,'Sheet 10 — Bulk powder Raman intensity of the reference sets (400–1800 cm-1)',['Raw counts include dark (~1200 counts) and background; marker intensity is baseline-corrected peak height.'])
sheet(wb,'11_Future_Acquisition',pd.DataFrame(FUTURE,columns=['Biomarker','Starting condition','Notes']),'Sheet 11 — Future measurement quick reference',widths={'Notes':90})
RC=pickle.load(open(os.path.expanduser('~/work/reaudit_cls.pkl'),'rb'))
sheet(wb,'12_NAM_60s_ReAudit',RC,'Sheet 12 — Individual re-audit of the 13 NAM 60 s × 5 acc spectra (A valid / B saturated / C distorted / D other)',['No spectrum is saturated inside 400–1800 cm-1 (class B = 0). Saturation occurs only at the laser line (≈ −2 to −4 cm-1).'],widths={'reason':90})
O,S70,CR=pickle.load(open(os.path.expanduser('~/work/nam70.pkl'),'rb'))
S70=S70.round(3); S70.columns=['Group','n','SNR mean','SNR SD','Mean pairwise r','Mean r to reference mean (glass zone excl.)','Strongest marker height (counts)','SD','Baseline (counts)']
O2=O[O.peak.isin([458,541,830,870,929,955,1023,1127])].copy().round(2); O2.columns=['Group','Peak (nominal)','n','Position mean','Position SD','Intensity mean','Intensity SD','CV (%)','SNR','Prominence','Prominence/noise','Detected','Relative intensity','Glass zone']
ws=sheet(wb,'13_NAM_70pct_vs_Reference',O2,'Sheet 13 — NAM 70 %/60 s/5 acc (n = 3, supporting group) vs reference 90 %/30 s/5 acc (n = 15)',
 ['Decision: 70 % spectra are ADDITIONAL SUPPORTING NAM spectra forming a separate acquisition-condition group; NOT merged into the reference (different power and integration time; n = 3; single session). Positions agree within ≤ 1 cm-1; r to reference = 0.998.'])
r0=ws.max_row+2
for j,c in enumerate(S70.columns,1): ws.cell(r0,j,c).font=H; ws.cell(r0,j).fill=HF
for i,row in enumerate(S70.itertuples(index=False),r0+1):
  for j,v in enumerate(row,1): ws.cell(i,j,v.item() if isinstance(v,np.generic) else v).font=F
CR2=CR.copy(); CR2.columns=['Peak (nominal)']+[c.replace('det_','Detected % ') for c in CR.columns[1:-1]]+['Mean detection %']
sheet(wb,'14_NAM_CrossDay_Detection',CR2,'Sheet 14 — NAM peak detection frequency across measurement days (all Valid NAM spectra; prominence ≥ 5 × noise)',['Used for Criterion 1 (reproducibility) of the NAM peak hierarchy.'])
wb.save(os.path.join(LIB,'05_Tables','Spectral_Library_Metadata_and_Peaks.xlsx'))
Lt.to_csv(os.path.join(LIB,'06_Literature','Literature_References.csv'),index=False); Lb.to_csv(os.path.join(LIB,'06_Literature','Literature_Band_List.csv'),index=False); M4.to_csv(os.path.join(LIB,'06_Literature','Experimental_vs_Literature_Peak_Assignments.csv'),index=False)
with open(os.path.join(LIB,'06_Literature','References.txt'),'w') as fh:
  for r in REFS: fh.write(f"[{r['id']}] {r['cite']} https://doi.org/{r['doi']}\n")
  fh.write('\nSearch note (23 Sep 2026): searches for N-acetylmuramic acid / MurNAc Raman & SERS, peptidoglycan Raman, lipoteichoic acid Raman & SERS, LPS / endotoxin / lipid A Raman did not locate a peer-reviewed normal-Raman spectrum of pure NAM or of isolated LTA powder. Two PMC full texts (PMC9575096, PMC10419157) could not be accessed (rate limit).\n')
Q.to_csv(os.path.join(LIB,'07_QC_and_Exclusions','QC_Log_All_Spectra.csv'),index=False); E.to_csv(os.path.join(LIB,'07_QC_and_Exclusions','Excluded_and_Flagged_Spectra.csv'),index=False)
S=dict(pk=M4.to_dict('records'),pkraw=PK.to_dict('records'),rep=REP.to_dict('records'),bulk=BULK.to_dict('records'),cond=COND.to_dict('records'),refs=REFS,litbands=Lb.to_dict('records'),future=FUTURE,
 counts=B.groupby(['cls','qc_status']).size().reset_index().values.tolist(),nfiles=len(A),n_out=int((A.cls=='OUT_OF_SCOPE').sum()),
 members={a:list(B[B.in_reference&(B.cls==a)].new_name) for a in ['NAM','LTA']},glasszone=[float(X[GZ].min()),float(X[GZ].max())],
 folders=B.groupby('raw_folder').size().to_dict(),nam_all_r=float(B[(B.cls=='NAM')&(B.qc_status=='Valid')].r_class_exglass.mean()),
 nam_valid_dates=sorted(B[(B.cls=='NAM')&(B.qc_status=='Valid')].ymd.unique()),lps=B[B.cls=='LPS'][['power','int_s','acc','snr']].to_dict('records'))
json.dump(S,open(os.path.expanduser('~/mnt/sukesh/_to_delete/claude_scratch/summary_v2.json'),'w'),default=lambda o:o.item() if isinstance(o,np.generic) else str(o),indent=1)
print('ok'); print(M4[['Biomarker','Peak category','Mean ± SD (cm-1)','Mean intensity (counts, baseline-corr.)','SNR','Replicate frequency (detected / n)','Confidence']].to_string())
