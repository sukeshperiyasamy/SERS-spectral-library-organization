import os,pickle,numpy as np,pandas as pd,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from lib2 import LIB
D=pickle.load(open('lp4.pkl','rb'));df=D['df'];P=D['P'];x=D['x'];BT=D['BT'];core=D['core']
GZ=(x>=1308)&(x<=1455);keep=~GZ
def nz(c): c=c-c[keep].min(); return c/c[keep].max()
for _,r in df.iterrows():
  p=P[r.key]['p'];folder='LPS' if r.category=='LPS_Bulk' else 'Glass_Blank'
  out=os.path.join(LIB,'02_Processed_Spectra',folder,r.library_file.replace('.csv','_processed.csv'))
  if os.path.exists(out): continue
  hdr=[f'# Library name: {r.library_file}',f'# Original raw file (unchanged): 28-9-26-lps-bulk/{r.orig}',
   '# Sample: '+('LPS bulk powder on glass slide (session 28 Sep 2026)' if r.category=='LPS_Bulk' else 'Empty glass slide (matched blank), LPS session 28 Sep 2026'),
   f'# Acquired (instrument clock): {r.acq} | Instrument: B&W Tek i-Raman Plus BWS465-785H | 785 nm | Objective: not recorded in file',
   f'# Laser power level: {r.power} % (instrument setting; not mW) | Integration: {r.int_s:g} s | Accumulations: {r.acc}',
   f'# QC: {r.QC} {r.QC_reason} | noise {r.noise:.1f} | strongest band {r.strong_pos:.0f} cm-1, SNR {r.snr:.1f} | ripple {r.ripple:.1f} | dark median {r.dark_med:.0f} | saturated px 400-1800: {r.nsat}'+(f' | {r.note}' if r.note else ''),
   '# Pipeline v3: hot-pixel interpolation (CCD px 640, 700); crop 400-1800 cm-1; MAD spike removal (z>6, width<=2 px); Savitzky-Golay 11/3; ALS baseline lambda=1e5 p=0.01 10 it; min-max normalisation',
   '# Columns: Raw; DarkSubtracted (= Raw - Dark); HotPixelFixed; Despiked; Smoothed; Baseline (ALS); BaselineCorrected (counts); Normalised (0-1)']
  t=pd.DataFrame({'Raman_Shift_cm-1':p['x'],'Raw_counts':p['raw'],'DarkSubtracted_counts':p['ds'],'HotPixelFixed_counts':p['ds_hotfix'],'Despiked_counts':p['despiked'],'Smoothed_counts':p['smooth'],'Baseline_counts':p['base'],'BaselineCorrected_counts':p['corr'],'Normalised_0to1':p['norm']}).round(5)
  with open(out,'w') as fh: fh.write('\n'.join(hdr)+'\n'); t.to_csv(fh,index=False)
qc=df[['library_file','category','power','int_s','acc','orig','md5','copy_ok','acq','dark_med','ripple','raw_max','nsat','sat_from','sat_to','ds_med','noise','strong_pos','strong_h','snr','rNAM','rLTA','rOLD','rBlank','r_to_others_mean','in_core','QC','QC_reason','note']].round(3)
qc.to_csv(os.path.join(LIB,'07_QC_and_Exclusions','LPS_Session_20260928_QC_Log.csv'),index=False)
ex=df[(df.category=='LPS_Bulk')&(df.QC.isin(['FAIL','REVIEW']))]
DONE_APPEND=True
if not DONE_APPEND: pd.DataFrame({'a':ex.library_file,'b':'28-9-26-lps-bulk/'+ex.orig,'c':'LPS','d':ex.power,'e':ex.int_s,'f':ex.acc,'g':np.where(ex.QC=='FAIL','Excluded (QC)','Review - not used in mean'),'h':ex.QC_reason,'i':'LPS_Bulk'}).to_csv(os.path.join(LIB,'07_QC_and_Exclusions','Excluded_and_Flagged_Spectra.csv'),mode='a',header=False,index=False)
# screening mean spectra (NOT reference)
RD=os.path.join(LIB,'03_Reference_Spectra','LPS')
for pw in (60,70):
  ks=[k for k in core if k.startswith(str(pw))];C=np.array([P[k]['p']['corr'] for k in ks]);Nm=np.array([nz(c) for c in C])
  pd.DataFrame({'Raman_Shift_cm-1':x,'BaselineCorrected_mean_counts':C.mean(0),'BaselineCorrected_median_counts':np.median(C,0),'BaselineCorrected_SD_counts':C.std(0,ddof=1),'Normalised_mean':Nm.mean(0),'Normalised_median':np.median(Nm,0),'Normalised_SD':Nm.std(0,ddof=1)}).round(5).to_csv(os.path.join(RD,f'LPS_Screening_20260928_P{pw:03d}pct_T30s_A05_Mean_Median_SD_NOT_REFERENCE.csv'),index=False)
BT.round(2).to_csv(os.path.join(RD,'LPS_Screening_20260928_Candidate_Bands.csv'),index=False)
# power summary
PS=[]
for pw in (60,70,80):
  s=df[(df.power==pw)&(df.category=='LPS_Bulk')];ks=[k for k in core if k.startswith(str(pw))]
  row=dict(power=pw,n_spectra=len(s),n_pass=int((s.QC=='PASS').sum()),n_review=int((s.QC=='REVIEW').sum()),n_fail=int((s.QC=='FAIL').sum()),n_core=len(ks),n_blanks=int(((df.power==pw)&(df.category=='Glass_Blank')).sum()))
  if len(ks)>1:
    Nm=np.array([nz(P[k]['p']['corr']) for k in ks]);rr=np.corrcoef(Nm[:,keep]);iu=np.triu_indices(len(ks),1)
    h=np.array([P[k]['p']['corr'][np.argmin(abs(x-425.4))-2:np.argmin(abs(x-425.4))+3].max() for k in ks])
    row.update(pairwise_r=rr[iu].mean(),r_mean_median=np.corrcoef(Nm.mean(0),np.median(Nm,0))[0,1],h425_mean=h.mean(),h425_cv=100*h.std(ddof=1)/h.mean(),snr425=np.mean([hh/P[k]['p']['noise'] for hh,k in zip(h,ks)]),r_mean_vs_April_LPS=np.corrcoef(Nm.mean(0)[keep],D['OLD'][keep])[0,1],r_mean_vs_NAM=np.corrcoef(Nm.mean(0)[keep],D['NAM'][keep])[0,1])
  PS.append(row)
PS=pd.DataFrame(PS);print(PS.round(3).to_string())
dec=pd.DataFrame([
 ['Sample identity','Spectra do not match NAM (r 0.14-0.19) and match the April 2026 LPS data (r 0.75-0.82): same pattern with a sharp band at 425 cm-1.','Accepted as LPS (as labelled).'],
 ['Best power',f'70 %: highest reproducibility (pairwise r {PS.pairwise_r[1]:.3f}), 425 band CV {PS.h425_cv[1]:.0f} %, SNR {PS.snr425[1]:.0f}. 60 %: more spots but CV {PS.h425_cv[0]:.0f} %. 80 %: 0/5 pass (fixed-pattern ripple), as for NAM at 80-90 %.','Use 70 %. Do not use 80 or 90 %.'],
 ['Signal strength','Only one strong band (425 cm-1, SNR 20-25). Other bands (765, 997, 1081, 1126, 860-900) have SNR 4-10 per spectrum; glass band dominates 1308-1455.','Increase SNR: 70 % / 30 s / 10 acc (same frame, more averaging, SNR x1.4); thicker packed powder.'],
 ['Reference status','Core set at 70 % n = 4 (+1 glass-dominated); 60 % n = 8.','Not enough for a final LPS reference: collect 10-15 spots at 70 % / 30 s / 10 acc + 3 matched blanks + 1 short blank.']],columns=['Item','Finding','Decision'])
wbp=os.path.join(LIB,'05_Tables','Spectral_Library_Metadata_and_Peaks.xlsx')
with pd.ExcelWriter(wbp,engine='openpyxl',mode='a',if_sheet_exists='replace') as w:
  for n,t in [('33_LPS_20260928_QC',qc),('34_LPS_20260928_Power',PS.round(3)),('35_LPS_20260928_Bands',BT.round(2)),('36_LPS_20260928_Decision',dec)]: t.to_excel(w,sheet_name=n,index=False)
wb=openpyxl.load_workbook(wbp)
for n in wb.sheetnames[-4:]:
  ws=wb[n]
  for c in ws[1]: c.font=Font(bold=True);c.fill=PatternFill('solid',fgColor='D9D9D9');c.alignment=Alignment(wrap_text=True,vertical='top')
  for col in ws.columns:
    L=max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:60]);ws.column_dimensions[col[0].column_letter].width=min(max(10,L*0.9),70)
wb.save(wbp)
with open(os.path.join(RD,'README_LPS_reference_pending.txt'),'a') as fh: fh.write('''

UPDATE 28 Sep 2026 - LPS screening session (folder 28-9-26-lps-bulk; 60/70/80 % x 30 s x 5 acc; 19 LPS spectra + 10 blanks)
- QC: 60 %: 8 core + 1 glass-dominated; 70 %: 4 core + 1 glass-dominated; 80 %: 0/5 (fixed-pattern ripple). 90 % not measured.
- Identity: matches April 2026 LPS data (r 0.75-0.82), not NAM (r < 0.2). Candidate primary band 425 cm-1 (raw-data check confirms; not in blanks).
- Weak candidate bands 765, 997, 1081 (SNR 4-10); 1249 possibly glass; 1308-1455 glass-dominated.
- Best condition: 70 %. Reference still PENDING: collect 10-15 spots at 70 % / 30 s / 10 acc + matched blanks.
- Screening means saved here as *_NOT_REFERENCE.csv; bands in LPS_Screening_20260928_Candidate_Bands.csv; workbook sheets 33-36.
''')
print('ok')
