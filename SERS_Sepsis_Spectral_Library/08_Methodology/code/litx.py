import sys,os,pickle,numpy as np,pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
sys.path.insert(0,os.path.expanduser('~/work')); from lib2 import *; from lit_v2 import *
PK,ALL,REP,BULK,COND,res=pickle.load(open(os.path.expanduser('~/work/an2.pkl'),'rb'))
DOI={r['id']:r['doi'] for r in REFS}
EVID={'L1':'SERS','L2':'SERS','L4':'SERS','L3':'Biological-sample Raman','L5':'Band compilation (biological tissues)','L6':'Pure related compounds (carbohydrate review)','L7':'Pure related compounds (lipid review)','L8':'Biological-sample FTIR','L9':'NIR/DFT (not Raman)','L10':'Normal Raman, related pure polymers','L11':'UV resonance Raman'}
ev=lambda cit:'; '.join(EVID[c.strip()] for c in cit.replace('(+L8 IR)','').split(',') if c.strip() in EVID) or '—'
rows=[]
for _,p in PK.iterrows():
  cat,lp,diff,ass,src,tech,cit,conf=MLIT[(p.analyte,p.nominal)]
  dois='; '.join(DOI[c.strip()] for c in cit.split(',') if c.strip() in DOI) if cit!='—' else '—'
  note=NOPURE if cat!='Glass/background' else 'Glass-associated feature (empty glass slide, this work)'
  rows.append([p.analyte,cat,f"{p.pos_mean:.1f}",lp,diff,ass,src,tech,ev(cit),cit,dois,conf,note])
# LPS rows
for bm,rid,tech,band,ass in LITBANDS:
  if bm!='LPS': continue
  rows.append(['LPS','Literature target (no experimental peak)','None confirmed',str(band),'—',ass,{'L1':'Commercial LPS/LOS (N. meningitidis, enteric)','L2':'Commercial endotoxins','L3':'Whole cells (A. baumannii, E. coli)'}[rid],tech,EVID[rid],rid,DOI[rid],'Uncertain',
     'No LPS reference yet; '+('band lies in glass zone 1308–1455 cm-1 – needs matched blank' if isinstance(band,(int,float)) and 1308<=band<=1455 else 'check in controlled 70/80/90 % experiment')])
for bm,rid,tech,band,ass in LITBANDS:
  if bm=='LTA':
    rows.append(['LTA','Literature SERS band (comparison)','no bulk match',str(band),'no experimental band within ±12 cm-1',ass,'Commercial LTA (B. subtilis)',tech,EVID[rid],rid,DOI[rid],'Low','SERS band; SERS and bulk normal Raman differ (surface selection rules)'])
D=pd.DataFrame(rows,columns=['Biomarker','Peak category','Experimental peak (cm-1)','Literature peak / region (cm-1)','Difference','Assignment (proposed)','Source/sample','Technique','Evidence type','Citation','DOI','Confidence','Note'])
wb=Workbook(); ws=wb.active; ws.title='Peak_Assignments'
H=Font(name='Arial',bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F4E79'); F=Font(name='Arial',size=10)
ws.cell(1,1,'Experimental bulk Raman bands vs literature — NAM, LTA, LPS (bulk powders on glass slides, 785 nm)').font=Font(name='Arial',bold=True,size=12)
ws.cell(2,1,'Key gap: '+NOPURE+' (NAM: none found; LTA: none found; LPS: SERS and whole-cell Raman only.)').font=Font(name='Arial',italic=True,size=10,color='C00000')
ws.cell(4,1,'Evidence types are kept separate: pure-compound Raman (none found for NAM/LTA/LPS) vs related pure compounds (reviews) vs biological-sample Raman/FTIR vs SERS. They are not interchangeable.').font=Font(name='Arial',italic=True,size=9)
ws.cell(3,1,'Confidence: High = compound-specific normal-Raman match; Moderate = well-established group frequency consistent with known structure plus related compound-class evidence; Low = general group-frequency region only; Uncertain = no assignment supported.').font=Font(name='Arial',italic=True,size=9)
for j,c in enumerate(D.columns,1): x=ws.cell(6,j,c); x.font=H; x.fill=HF; x.alignment=Alignment(wrap_text=True)
for i,row in enumerate(D.itertuples(index=False),7):
  for j,v in enumerate(row,1): c=ws.cell(i,j,v); c.font=F; c.alignment=Alignment(wrap_text=True,vertical='top')
for j,w in enumerate([9,22,12,18,14,42,30,24,28,10,30,11,45],1): ws.column_dimensions[get_column_letter(j)].width=w
ws.freeze_panes='C7'; ws.auto_filter.ref=f"A6:M{6+len(D)}"
ws2=wb.create_sheet('References'); ws2.column_dimensions['A'].width=8; ws2.column_dimensions['B'].width=120; ws2.column_dimensions['C'].width=30; ws2.column_dimensions['D'].width=60
for j,c in enumerate(['ID','Citation','DOI','Access / relevance note'],1): x=ws2.cell(1,j,c); x.font=H; x.fill=HF
for i,r in enumerate(REFS,2):
  for j,v in enumerate([r['id'],r['cite'],'https://doi.org/'+r['doi'],r['technique']+' | '+r['verified']],1): ws2.cell(i,j,v).font=F
ws3=wb.create_sheet('Search_Log'); ws3.column_dimensions['A'].width=60; ws3.column_dimensions['B'].width=90
log=[('N-acetylmuramic acid Raman / vibrational spectroscopy / MurNAc Raman / NAM SERS / peptidoglycan Raman NAM','No pure-compound normal Raman or SERS spectrum found. Related: NIR + DFT-calculated IR of NAM (L9, no band list); commercial IR spectrum entries (not peer-reviewed).'),
 ('lipoteichoic acid Raman / vibrational spectroscopy / SERS / S. aureus LTA Raman','No pure-compound normal Raman spectrum found. SERS of LTA (L2: 894, 1333 cm-1; L4: bands not accessible). FTIR of wall teichoic acid in whole S. aureus cells (L8).'),
 ('lipopolysaccharide Raman / SERS / LPS vibrational spectroscopy / E. coli LPS Raman / lipid A Raman','SERS band lists (L1, L2); whole-cell Raman of lipid A modification (L3). No bulk LPS powder normal-Raman band table found.'),
 ('Not accessible during search','PMC9575096 and PMC10419157 full texts (Europe PMC rate limit).')]
for j,c in enumerate(['Search terms','Outcome'],1): x=ws3.cell(1,j,c); x.font=H; x.fill=HF
for i,(a,b) in enumerate(log,2): ws3.cell(i,1,a).font=F; ws3.cell(i,2,b).font=F; ws3.cell(i,2).alignment=Alignment(wrap_text=True)
wb.save(os.path.join(LIB,'06_Literature','Literature_Peak_Assignments.xlsx')); print(len(D))
with open(os.path.join(LIB,'06_Literature','References.txt'),'w') as fh:
  for r in REFS: fh.write(f"[{r['id']}] {r['cite']} https://doi.org/{r['doi']}\n")
  fh.write('\n'+NOPURE+'\nSearch covered NAM/MurNAc, peptidoglycan, LTA and LPS/lipid A Raman, SERS and vibrational spectroscopy (23 Sep 2026). PMC9575096 and PMC10419157 full texts could not be accessed (rate limit).\n')
