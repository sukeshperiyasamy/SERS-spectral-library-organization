# Rewrites the LTA parts of the DFT comparison (v1.3) and simulation (v1.2) reports with the whole-molecule LTA model
# (simulation/lta output simulation, 28 Oct 2026) and the pooled LTA reference v2.1. Edits the .docx in place; backup in ../_backup_before_lta.
# Re-running restores the backed-up originals first.
import os,re,copy,json,shutil,io
from docx import Document
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
L=os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)),'..','..')); REP=os.path.join(L,'09_Reports'); FIG=os.path.join(L,'04_Figures','LTA_Model_20261028')
SIMF=os.path.join(L,'08_Methodology','simulation','LTA_Model_20261028','quick_B3LYP-3-21G')
BK=os.path.join(L,'_backup_before_lta'); os.makedirs(BK,exist_ok=True)
F={'dft':'DFT_vs_Experimental_Comparison_NAM_LPS_LTA.docx','sim':'Simulation_Methods_and_Spectra_NAM_LPS_LTA.docx'}
for f in F.values():
    if not os.path.exists(os.path.join(BK,f)): shutil.copy(os.path.join(REP,f),os.path.join(BK,f))
    else: shutil.copy(os.path.join(BK,f),os.path.join(REP,f))
R=json.load(open(os.path.join(L,'05_Tables','LTAmodel_vs_LTA_results.json'))); M=R['model']; FR=R['frag']; NB=R['n_bands']
def settext(p,text):
    rs=p.runs
    if rs:
        rs[0].text=text
        for r in rs[1:]: r._r.getparent().remove(r._r)
    else: p.add_run(text)
def rep(p,old,new):
    for r in p.runs:
        if old in r.text: r.text=r.text.replace(old,new); return True
    if old in p.text: settext(p,p.text.replace(old,new)); return True
    return False
def find_p(doc,start,style=None):
    for p in doc.paragraphs:
        if p.text.strip().startswith(start) and (style is None or p.style.name.startswith(style)): return p
    raise KeyError(start)
def cellset(c,text):
    p=c.paragraphs[0]
    if p.runs: settext(p,text)
    else: p.add_run(text)
    for extra in c.paragraphs[1:]: extra._p.getparent().remove(extra._p)
def addrow(t,vals):
    new=copy.deepcopy(t.rows[-1]._tr); t._tbl.append(new); row=t.rows[-1]
    for c,v in zip(row.cells,vals): cellset(c,v)
    return row
def delrow(t,row): t._tbl.remove(row._tr)
def shade(c,hx):
    pr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),hx); pr.append(sh)
class Ins:
    def __init__(s,doc,anchor): s.d=doc; s.cur=anchor if not hasattr(anchor,'_p') else anchor._p
    def _put(s,el): s.cur.addnext(el); s.cur=el
    def h(s,text,lvl=2): p=s.d.add_heading(text,lvl); s._put(p._p)
    def p(s,text,italic=False,size=None):
        p=s.d.add_paragraph(); r=p.add_run(text); r.italic=italic
        if size: r.font.size=Pt(size)
        s._put(p._p); return p
    def b(s,text): p=s.d.add_paragraph(text,style='List Bullet'); s._put(p._p)
    def fig(s,path,cap,w=17):
        s.d.add_picture(path,width=Cm(w)); p=s.d.paragraphs[-1]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; s._put(p._p)
        c=s.d.add_paragraph(); r=c.add_run(cap); r.italic=True; r.font.size=Pt(8.5); s._put(c._p)
    def t(s,head,rows,widths,fs=8):
        t=s.d.add_table(rows=1,cols=len(head)); t.style='Table Grid'
        for i,x in enumerate(head):
            c=t.rows[0].cells[i]; c.text=''; r=c.paragraphs[0].add_run(x); r.bold=True; r.font.size=Pt(fs); shade(c,'D9D9D9')
        for row in rows:
            cs=t.add_row().cells
            for i,x in enumerate(row): cs[i].text=''; cs[i].paragraphs[0].add_run(str(x)).font.size=Pt(fs)
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
        s._put(t._tbl); sp=s.d.add_paragraph(); s._put(sp._p)
def tb(d,*heads):
    for t in d.tables:
        if [c.text for c in t.rows[0].cells][:len(heads)]==list(heads): return t
    raise KeyError(heads)
def f2(x): return f'{x:.2f}'
def sgn(x): return f'{x:+.1f}'
def pstr(p): return '< 0.001' if p<0.0005 else f'{p:.2f}'
def tchance(blk,t): b=blk[t]; return f"{b['decoy_mean']:.1f} +/- {b['decoy_sd']:.1f}",pstr(b['p'])
CL={b[0]:b[1] for b in R['bands_used']}
def swap_image(d,cap_start,png):
    cap=find_p(d,cap_start); prev=cap._p.getprevious()
    rid=prev.xpath('.//a:blip/@r:embed')[0]; part=d.part.related_parts[rid]; part._blob=open(png,'rb').read()
# ================================================================= DFT comparison
def dft():
    d=Document(os.path.join(REP,F['dft']))
    rep(d.paragraphs[2],'Version 1.2, 28 October 2026','Version 1.3, 28 October 2026')
    rep(d.paragraphs[2],'(Gaussian 16W, 27-28 Oct 2026).','(Gaussian 16W, 27-28 Oct 2026); LTA whole-molecule model: SERS_Sepsis_Spectral_Library/08_Methodology/simulation/LTA_Model_20261028 (Gaussian 16W, 28 Oct 2026). LTA experimental reference: pooled v2.1 (n = 17) of 28 Oct 2026.')
    t5,t10,t18=M['t5'],M['t10'],M['t18']
    settext(d.paragraphs[3],"The simulations reproduce the position of many individual bands, and the agreement is best for NAM, the only small molecule calculated as a whole. LPS and LTA each now have two calculations: separate fragments, and a whole-molecule model (shortened Kdo2-lipid A, 201 atoms; three glycerol-phosphate units with two D-alanine esters and one alpha-GlcNAc, 93 atoms). For LPS the model puts all five library bands within 2.4 cm-1 of a calculated mode, but this is no better than chance. For LTA the model puts all nine primary and supporting bands of the new pooled reference (v2.1, 17 spots) within 4.5 cm-1 of a calculated mode (mean difference "+f"{M['mae']:.1f}"+" cm-1, against "+f"{FR['mae']:.1f}"+" cm-1 for the three fragments) and, at the strict 5 cm-1 tolerance, this is above chance (9 of 9 against "+tchance(M,'t5')[0]+" expected, p "+pstr(t5['p'])+"); at 10 and 18 cm-1 it is not (p = "+f"{t10['p']:.2f} and {t18['p']:.2f}"+"). The LTA modes that sit on the library bands are weak (relative activity 0.01 to 0.15), the strongest calculated modes (690, 654 and 1490-1510 cm-1) have no counterpart in the measured spectrum, and the whole-spectrum correlation is zero ("+f"{M['r_whole']:.2f}"+"). The overall spectral shapes do not match for any molecule (correlation 0.15 or less). The DFT results can support tentative assignments of individual bands, but they cannot be used to confirm identity, and several strong experimental bands remain unexplained: NAM 829 and 955 cm-1, the assignment of LPS 425 cm-1, and for LTA the strong 1092 and 1010 cm-1 bands, which match only very weak modes.")
    rep(find_p(d,'The bands tested are the primary'),'and for none of the three molecules does it','and only the LTA whole-molecule model does, and only at the 5 cm-1 tolerance')
    # --- summary table
    t=tb(d,'Molecule','What was simulated')
    for r in t.rows:
        if r.cells[0].text=='LTA':
            c5,c10=tchance(FR,'t10')
            for c,v in zip(r.cells,['LTA (3 fragments)','Three separate fragments, not whole LTA',str(NB),f"{FR['t5']['real']} / {FR['t10']['real']} / {FR['t18']['real']}",f"{FR['mae']:.1f}",f"{FR['t10']['decoy_mean']:.1f} (p = {FR['t10']['p']:.2f})",f2(FR['r_whole']),'Weak']): cellset(c,v)
    addrow(t,['LTA (whole-molecule model)','3 glycerol-phosphates, 2 D-Ala esters, 1 alpha-GlcNAc; 93 atoms, net charge 0, B3LYP/3-21G, gas phase',str(NB),f"{t5['real']} / {t10['real']} / {t18['real']}",f"{M['mae']:.1f}",f"{t10['decoy_mean']:.1f} (p = {t10['p']:.2f})",f2(M['r_whole']),'Positions agree, above chance at 5 cm-1 only; matched modes weak; shape no'])
    # --- what was simulated table
    t=tb(d,'Item','NAM','LPS','LTA')
    add={'What was calculated':' Also: whole-molecule model (built, no PubChem entry): three glycerol-phosphate units joined by two phosphodiesters, D-alanine ester on each outer glycerol, alpha-GlcNAc on the central glycerol; C23H45N3O20P2, 93 atoms, net charge 0 (two phosphate anions, two ammonium groups).',
      'Why fragments':'. For the whole-molecule model the chain was kept to three units and the glycolipid anchor left out so that the molecule stays computable; about 23 repeats are present in the real polymer.',
      'Method':' Whole-molecule model: Gaussian 16W Rev C.01, PM6 pre-optimisation then B3LYP/3-21G opt freq=raman; 75 optimisation steps; 273 modes, no imaginary frequency (lowest 31 cm-1).',
      'Environment':'; whole-molecule model: gas phase, no counter-ions, no water',
      'Frequency scaling':'; whole-molecule model 0.965 (typical B3LYP value, not fitted for 3-21G)',
      'Modes in 400-1800 cm-1':f"; whole-molecule model {M['n_modes']}",
      'Modes per 100 cm-1':f"; whole-molecule model {M['n_modes']/14:.1f}"}
    for r in t.rows:
        k=r.cells[0].text
        if k in add: cellset(r.cells[3],r.cells[3].text+add[k])
        if k=='Experimental sample': cellset(r.cells[3],'Sigma L2515-10MG, S. aureus LTA, lot 0000419622; library reference v2.1 (pooled n = 17, 28 Oct 2026). The earlier fragment comparison used the preliminary April 2026 reference (n = 4).')
    p13=find_p(d,'A calculation on a single small fragment')
    settext(p13,"A calculation on a single small fragment cannot reproduce a polymer. LPS has phosphorylated, acylated and branched sugars; LTA has a charged phosphodiester backbone. The five LPS fragments and the three LTA fragments contain none of the linkages between units, so bands that come from those linkages cannot appear in them. The whole-molecule LPS model does contain the glycosidic bond between the two glucosamines, the Kdo links, both phosphate esters and both amide links of the lipid A backbone, but not the long acyl chains, the outer core sugars or the O-antigen. The whole-molecule LTA model contains two phosphodiester links, two D-alanine ester links and the glycosidic link to GlcNAc, but only three repeating units instead of about 23 and no glycolipid anchor. Both are therefore parts of the molecule and not the whole.")
    # --- section 6 rebuild
    h6=find_p(d,'6. LTA','Heading 1'); h7=find_p(d,'7. Is the agreement better than chance?','Heading 1')
    els=list(d.element.body.iterchildren()); a,b=els.index(h6._p),els.index(h7._p)
    for el in els[a+1:b]: d.element.body.remove(el)
    I=Ins(d,h6)
    I.h('6.1 Result',2)
    I.p(f"The LTA reference used here is the pooled v2.1 reference of 28 October 2026 (17 independent spots, clean dark, 80 % power; 10 spots at 45 s x 2 acc and 7 at 30 s x 3 acc). Nine primary and supporting bands outside the glass zone were tested (primary 539, 862 and 1092 cm-1). The whole-molecule model puts all nine within 5 cm-1 of a calculated mode (mean absolute difference {M['mae']:.1f} cm-1). The three fragments reach {FR['t5']['real']}, {FR['t10']['real']} and {FR['t18']['real']} of nine within 5, 10 and 18 cm-1 (mean difference {FR['mae']:.1f} cm-1).")
    I.p(f"Chance test. If the model modes are shifted by a random 15 to 70 cm-1, {tchance(M,'t5')[0]} bands still fall within 5 cm-1 by chance; the real model gets 9 of 9 (none of 5000 random shifts did as well, p {pstr(t5['p'])}). At 10 cm-1 chance gives {tchance(M,'t10')[0]} and the model is not clearly better (p = {t10['p']:.2f}); at 18 cm-1 it is not (p = {t18['p']:.2f}). The fragments are indistinguishable from chance at all three tolerances (p = {FR['t5']['p']:.2f}, {FR['t10']['p']:.2f}, {FR['t18']['p']:.2f}). The model has {M['n_modes']} modes in 400-1800 cm-1 (11.8 per 100 cm-1), the fragments only {FR['n_modes']} peaks (3.4 per 100 cm-1), so the model has a much higher baseline.")
    I.p(f"What the agreement does not show. The calculated modes that sit on the library bands are weak. Their relative activities are 0.01 to 0.15 on a scale where the strongest calculated mode is 1.0, so counting only strong modes (activity 0.15 or more) the model matches {M['t5']['real_strong']} of 9 bands within 5 cm-1 and is not above chance (p = {M['t5']['p_strong']:.2f}). The whole-spectrum correlation with the reference is {M['r_whole']:.2f} for the model and {FR['r_whole']:.2f} for the fragments. Fitting weights to the fragments and the model together (NNLS) gives r = {R['nnls_all']['r']:.2f}, with the model taking {R['nnls_all']['weights']['LTA model']*100:.0f} % of the weight; the fragments alone give r = {FR['nnls_r']:.2f}.")
    I.fig(os.path.join(FIG,'LTA_Exp_vs_Model.png'),"Figure 5. LTA. Top: library reference v2.1 (17 spots), dashed lines mark the nine primary and supporting bands (green: nearest calculated mode within 10 cm-1). Bottom: whole-molecule model, calculated modes (blue sticks, scaled 0.965) and broadened spectrum (black). The model's strongest mode, 690 cm-1, falls where the measured spectrum has a minimum.")
    I.fig(os.path.join(FIG,'LTA_Exp_vs_Fragments.png'),"Figure 6. LTA, fragment calculation against the same v2.1 reference: broadened DFT curves of the three fragments (scaled 0.96) and their equal-weight average.")
    I.h('6.2 Band-by-band',2)
    rows=[]
    for m,f in zip(M['bands'],FR['bands']):
        rows.append([f"{m['exp']:.1f}",CL[m['exp']],f"{m['calc']:.1f}",sgn(m['delta']),f"{m['act']:.2f}",f"{f['calc']:.1f}",sgn(f['delta']),f"{f['act']:.2f}"])
    I.t(['Exp. band (cm-1)','Class','Model: nearest mode','Exp - calc','Rel. activity','Fragments: nearest mode','Exp - calc','Rel. activity (fragment)'],rows,[1.9,1.9,2.2,1.7,1.7,2.4,1.7,2.3])
    I.fig(os.path.join(FIG,'LTA_Model_vs_Fragments.png'),"Figure 7. LTA: whole-molecule model against the three fragments. Left: difference between each library band and the nearest calculated mode. Middle: relative activity of that mode. Right: whole-spectrum correlation with the reference.")
    I.h('6.3 What is happening and why',2)
    I.b("Positions. Every library band has a model mode within 4.5 cm-1, including the two that the fragments missed by more than 30 cm-1 (799 cm-1 against 768; 1751 cm-1 against 1690) and the 1092 cm-1 primary band (model 1089.5; nearest fragment peak 14 cm-1 away). Part of the reason is that the model contains the phosphodiester and ester links. Part is density: the model has 165 modes in the window, so a mode is rarely more than a few cm-1 from any band.")
    I.b("Intensities. The 1092 cm-1 band is the strongest measured band (SNR 30), but the model mode near it has a relative activity of only 0.03; the strong measured 799 and 1010 cm-1 bands match activities of 0.08 and 0.01. The model's strongest modes are 690 cm-1 (1.00), 654, 594, 446 and a group at 1490-1510 cm-1 (0.2 to 0.3). The measured spectrum has no band at 690 (a minimum) and only weak features at 1583 and 1459 cm-1. The 1490-1510 group is probably NH3+ and amide motion exaggerated by the small basis set and the lack of water; this was not checked against the displacement vectors.")
    I.b("Shape. The correlation is zero (-0.01) for the model and 0.09 for the fragments, so neither reproduces the measured pattern. Adding the model to the fragments in the NNLS fit does not improve the fit (r 0.08).")
    I.b("The 5 cm-1 result is therefore a modest finding: positions are consistent with a phosphate-ester, glycerol and D-alanine/GlcNAc model, which supports tentative assignments, but the modes are not the ones that dominate the measured spectrum and the 10 and 18 cm-1 tests are within chance.")
    I.h('6.4 Likely reasons for the differences',2)
    I.b("Model, not polymer. Three repeating units instead of about 23, no glycolipid anchor; chain-end effects are exaggerated and repeat-unit bands are under-represented. Glycerol stereochemistry of the two outer arms was set equal and not checked against the S. aureus form.")
    I.b("Low-level theory. B3LYP/3-21G has no polarisation or diffuse functions and is poor for charged phosphate and ammonium groups; the 0.965 scale factor was not fitted for it. Peaks near 2410, 2510 and 2700-2800 cm-1 are probably artefacts of intramolecular NH3+...phosphate hydrogen bonds (outside the window used here).")
    I.b("Gas phase, one conformer, no counter-ions. The powder contains Na+, Mg2+ or Ca2+ and probably water; the experimental bands are broad and sit on a fluorescence background.")
    I.b("Experiment. The reference is better than the April one (17 spots, clean darks), but the primary-band heights vary by about 26-35 % between spots, and normalisation to the maximum compresses the weaker bands.")
    # --- chance section
    rep(find_p(d,'Any list of calculated modes'),'Figure 6 and the table below show that it does not, for any of the three molecules, and the same holds for the LPS whole-molecule model.','Figure 8 and the table below show that it does not for NAM, the LPS fragments, the LPS whole-molecule model or the LTA fragments. The LTA whole-molecule model is the exception at the strict 5 cm-1 tolerance only.')
    swap_image(d,'Figure 6. Grey','chance') if False else None
    cp=find_p(d,'Figure 6. Grey'); settext(cp,'Figure 8. Grey: number of experimental bands within 10 cm-1 of a calculated mode after random shifts of the modes. Coloured line: real calculation. All five calculations are shown (LTA tested against reference v2.1).')
    swap_image(d,'Figure 8. Grey',os.path.join(FIG,'D_Chance_Test_all.png'))
    cp=find_p(d,'Figure 7. Difference'); settext(cp,'Figure 9. Difference between every library band and its nearest calculated mode, all five calculations. Points outside the green band (10 cm-1) are bands the calculation does not explain.')
    swap_image(d,'Figure 9. Difference',os.path.join(FIG,'E_Position_Differences_all.png'))
    # chance table
    t=tb(d,'Molecule','Tolerance (cm-1)')
    for r in list(t.rows):
        if r.cells[0].text=='LTA': delrow(t,r)
    for nm,blk in (('LTA (3 fragments)',FR),('LTA (whole model)',M)):
        for tol in (5,10,18):
            x=blk[f't{tol}']; addrow(t,[nm,str(tol),f"{x['real']} of {NB}",f"{x['decoy_mean']:.1f} +/- {x['decoy_sd']:.1f}",pstr(x['p']) if x['p']<0.0005 else f"{x['p']:.2f}",f"{x['real_strong']} of {NB} (p = {x['p_strong']:.2f})"])
    rep(find_p(d,'Here p is the fraction'),'All p values are 0.19 or higher, far from 0.05, and using only the strong modes does not change this.','All p values are 0.19 or higher except the LTA whole-molecule model at 5 cm-1 (p < 0.001) and at 10 cm-1 (p = 0.07); using only the strong modes, no calculation is above chance.')
    rep(find_p(d,'Here p is the fraction'),'the agreement is described as not demonstrated rather than good.','the agreement is described as not demonstrated rather than good, with the LTA model at 5 cm-1 as a limited exception.')
    # conclusion table
    t=tb(d,'Question','NAM','LPS','LTA')
    vals={'Do individual band positions agree?':'Yes: all 9 bands within 4.5 cm-1 (whole-molecule model); fragments: 5 of 9 within 10 cm-1, 1092 and 1751 off by 14 and 61 cm-1',
      'Is that better than chance?':f"Whole model: yes at 5 cm-1 (p {pstr(t5['p'])}), no at 10 and 18 cm-1 (p = {t10['p']:.2f}, {t18['p']:.2f}); fragments: no (p = {FR['t10']['p']:.2f})",
      'Does the overall spectral shape agree?':f"No (whole model r = {M['r_whole']:.2f}; fragments {FR['r_whole']:.2f}, fitted {FR['nnls_r']:.2f})",
      'Are the strongest experimental bands explained?':'Positions of 1092, 799, 862 and 539 are matched, but by weak modes (activity 0.01 to 0.15); the strongest model modes (690, 1490-1510) are not seen in the measurement',
      'What can be claimed?':'Position consistency with a phosphodiester / glycerol / D-alanine / GlcNAc model supports the existing tentative assignments (1092 phosphate PO2-, 862 and 799 glycerol-phosphate and sugar, 1459 CH2, 1751 ester C=O); no new assignment gained',
      'Overall':'POSITIONS AGREE (above chance at 5 cm-1 only); intensities and shape not matched; not demonstrated'}
    for r in t.rows:
        if r.cells[0].text in vals: cellset(r.cells[3],vals[r.cells[0].text])
    # conclusions / next steps / limitations
    rep(find_p(d,'LTA: model a repeating glycerol-phosphate'),'LTA: model a repeating glycerol-phosphate dimer or trimer with its counter-ion, with D-alanine attached as an ester; use the same larger basis set as NAM.','LTA: the three-unit model is a first whole-molecule result at a quick level. Next steps: finish the B3LYP/6-31+G(d,p) water (PCM) run that was started on 28 Oct; add counter-ions, a longer chain (4-5 units) and the Glc2-diacylglycerol anchor; check the glycerol stereochemistry; look at the atoms that move in the 1089, 862 and 802 cm-1 modes and in the strong 690 cm-1 mode, which the measurement does not show.')
    rep(find_p(d,'All: report the same chance test'),'and compare against the new LTA reference once 12 to 15 spots are acquired.','and the LTA comparison now uses the v2.1 reference (17 spots).')
    rep(find_p(d,'LTA peaks were picked'),'LTA peaks were picked from broadened curves (prominence 0.04), so weak modes may be missed.','LTA fragment peaks were picked from broadened curves (prominence 0.04), so weak modes may be missed; the whole-molecule model has a full mode list.')
    rep(find_p(d,'Only primary and supporting library bands'),'(NAM 11, LPS 5, LTA 4)','(NAM 11, LPS 5, LTA 9)')
    p=find_p(d,'The LTA comparison in Section 6 still uses'); settext(p,'The LTA comparison in Section 6 uses the pooled v2.1 reference of 28 Oct 2026 (17 spots, pending operator confirmation). The LTA whole-molecule model is a three-unit, gas-phase B3LYP/3-21G calculation; its band assignments were not checked against displacement vectors, the glycerol stereochemistry was not verified, and the scale factor 0.965 was not fitted for this basis set. The 5 cm-1 chance test depends on random shifts of 15 to 70 cm-1 and on the high mode density.')
    # file table
    t=tb(d,'Item','File used')
    rows=[r for r in t.rows if r.cells[0].text.startswith('LTA fragment curves')]
    if rows: cellset(rows[0].cells[1],rows[0].cells[1].text)
    addrow(t,['LTA whole-molecule model (165 modes in range)','SERS_Sepsis_Spectral_Library\\08_Methodology\\simulation\\LTA_Model_20261028 (copy of C:\\Users\\sukes\\Downloads\\sukesh\\simulation\\lta output simulation): lta_model_raman_act.csv, lta_model.gjf, lta_model.out (Gaussian log), LTA_Raman_Methods.md, LTA_Raman_Full_Report.docx'])
    addrow(t,['LTA reference used','03_Reference_Spectra\\LTA\\LTA_Bulk_Reference_20261028_Pooled_Mean_SD.csv and LTA_Reference_20261028_Pooled_Marker_Peaks.csv (v2.1, n = 17)'])
    addrow(t,['LTA comparison results and code','05_Tables\\LTAmodel_vs_LTA_results.json; 08_Methodology\\code\\lta_model_compare.py, lta_figures.py, update_docs_lta.py; figures in 04_Figures\\LTA_Model_20261028'])
    d.save(os.path.join(REP,F['dft']))
# ================================================================= Simulation methods
def sim():
    d=Document(os.path.join(REP,F['sim']))
    rep(d.paragraphs[2],'Version 1.1, 28 October 2026','Version 1.2, 28 October 2026')
    p=d.paragraphs[3]
    rep(p,'LTA was simulated as three separate fragments. The intact LPS and LTA molecules are too large and variable to calculate (Section 7), so no whole-LPS or whole-LTA spectrum was simulated.','LTA was simulated as three separate fragments and, on 28 Oct 2026, as a 93-atom whole-molecule model of three glycerol-phosphate units carrying two D-alanine esters and one alpha-GlcNAc. The intact LPS and LTA molecules are too large and variable to calculate (Section 7), so no spectrum of the intact LPS or LTA was simulated.')
    h6=find_p(d,'6. LTA','Heading 1'); settext(h6,'6. LTA: three-fragment simulation and whole-molecule model')
    t=tb(d,'','NAM','LPS','LTA')
    add={'What was simulated':' Also: one whole-molecule model (3 glycerol-phosphate units, 2 D-Ala esters, 1 alpha-GlcNAc).',
      'Number of atoms per calculation':'. Whole-molecule model: 93 (C23H45N3O20P2, net charge 0)',
      'Software':'. Whole-molecule model: Gaussian 16W Rev C.01',
      'Method and basis set':'. Whole-molecule model: PM6 pre-optimisation, then B3LYP / 3-21G',
      'Environment':'. Whole-molecule model: gas phase, no counter-ions',
      'Frequency scale factor':'. Whole-molecule model: 0.965 (not fitted for 3-21G)',
      'Imaginary frequencies':'. Whole-molecule model: none (273 modes; true minimum)',
      'Spectra obtained':'. Whole-molecule model: 273 modes (165 in 400-1800 cm-1), broadened 785 nm spectrum 200-3300 cm-1',
      'Why this approach':'. Whole-molecule model: a short piece of the polymer that carries all three substituent types'}
    for r in t.rows:
        k=r.cells[0].text
        if k in add: cellset(r.cells[3],r.cells[3].text+add[k])
    p73=find_p(d,'The three curves were averaged with equal weights')
    rep(p73,'are the earlier fit to April data and are reported here as found.','are the earlier fit to April data and are reported here as found; the comparison document now tests the fragments against the v2.1 reference.')
    I=Ins(d,p73)
    I.h('6.5 Whole-molecule model: three glycerol-phosphate units with D-alanine and GlcNAc',2)
    I.p("Why. The earlier fragment sum ignores the bonds that join the pieces, so a model that keeps them was built. There is no PubChem entry for LTA (a polymer), so the structure was built with RDKit 2026.03.3, not downloaded. S. aureus LTA is about 23 glycerol-phosphate repeats on a Glc2-diacylglycerol anchor; position 2 of each glycerol can be free, esterified with D-alanine or carry N-acetylglucosamine. The model carries all three in a three-unit chain.")
    I.t(['Item','Value'],[
      ['Composition','Three glycerol-phosphate units joined by two phosphodiester links; D-alanine ester on each of the two outer glycerols; alpha-GlcNAc on the central glycerol'],
      ['Formula, atoms, charge','C23H45N3O20P2; 93 atoms; net charge 0, singlet (two phosphodiesters -1, two alanine ammonium groups +1)'],
      ['Stereochemistry','GlcNAc from LTA\\N-Acetyl-D-Glucosamine.sdf (beta-anomer, R at C1) flipped to alpha (S at C1) because LTA GlcNAc is alpha-linked; D-alanine R (D) on both esters, checked with RDKit CIP labels. The two glycerol arms were set the same and not checked against the sn-glycerol-1-phosphate form'],
      ['SMILES','CC(=O)N[C@@H]1[C@@H](O)[C@H](O)[C@@H](CO)O[C@@H]1O[C@H](COP(=O)([O-])OC[C@H](OC(=O)[C@@H](C)[NH3+])CO)COP(=O)([O-])OC[C@H](OC(=O)[C@@H](C)[NH3+])CO'],
      ['Starting geometry','Best of 300 ETKDGv3 conformers (seed 7), each MMFF-optimised; lowest MMFF energy -311.4 kcal/mol'],
      ['Software','Gaussian 16W Rev C.01 (C:\\G16W\\g16.exe); RDKit; Python (numpy, scipy, matplotlib)'],
      ['Job 1','# opt pm6, gas phase, default convergence, charge 0 (finished 09:29)'],
      ['Job 2','# opt freq=raman b3lyp/3-21g geom=allcheck guess=read, gas phase; optimisation converged at step 75 (12:16); frequencies 49 min; normal termination 13:05'],
      ['Resources','%nprocshared=28, %mem=54GB (i7-14700 PC); job queued after the lipid A job']],[4,13])
    I.h('6.6 Processing of the LTA model calculation',2)
    I.t(['Step','Detail'],[
      ['Extraction','Harmonic frequencies and Raman activities from lta_model.out into lta_model_raman_act.csv (273 modes, unscaled; lowest 31 cm-1, none imaginary)'],
      ['Scaling','Frequencies x 0.965 (typical B3LYP value; not fitted or benchmarked for 3-21G)'],
      ['Intensity','I = S (v0 - v)^4 / [v (1 - exp(-hcv/kT))], v0 = 1e7/785 cm-1, T = 298.15 K'],
      ['Broadening','Lorentzian, FWHM 12 cm-1; 200-3300 cm-1; 6200 points; normalised to the tallest peak (make_spectrum.py)'],
      ['Comparison grid','For the comparison document the 165 modes in 400-1800 cm-1 were drawn as Gaussians of FWHM 12 cm-1 on a 1 cm-1 grid (lta_model_compare.py)']],[4,13])
    I.h('6.7 Spectrum obtained and its limits',2)
    I.p("Strongest scaled features in the fingerprint region (relative height): 689 (1.00), 1490 (0.80) and 1503 (0.66) in the broadened spectrum, then 816, 654, 1325, 1240, 914, 972 and 1006 (0.29 to 0.34). The 1490-1503 group is probably the NH3+ / amide region and 689 probably a phosphate P-O-C or ring deformation; these assignments were not verified against the vibrational displacements.")
    I.fig(os.path.join(SIMF,'lta_model_spectrum_785nm.png'),"Figure 6. Whole-molecule LTA model: calculated Raman spectrum at 785 nm (B3LYP/3-21G, scaled 0.965, FWHM 12 cm-1), from the calculation folder.")
    for x in ["Quick level: gas phase and a small basis set; positions are approximate and intensities distorted, especially for the charged phosphate and ammonium groups.",
              "Peaks near 2410, 2510 and 2700-2800 cm-1 are probably artefacts of strong NH3+...phosphate hydrogen bonds without water; ignore that region.",
              "Short chain (three units, no glycolipid anchor): chain-end effects exaggerated, repeat-unit bands under-represented. No counter-ions, no water, one conformer. Glycerol stereochemistry assumed.",
              "A higher-level run (B3LYP/6-31+G(d,p), water PCM, same level as NAM) was started from this geometry on 28 Oct 2026 and is not finished; its results are not in this document."]: I.b(x)
    rep(find_p(d,'The reasons below come from the project files'),'for LTA only fragments were run.','for LTA the three-unit model of Section 6.5 is the largest run so far.')
    rep(find_p(d,'Rerun the lipid A model at'),'and run a two-to-three-unit glycerol-phosphate oligomer with D-alanine ester and counter-ions for LTA, on a workstation with enough memory.','and finish the LTA model at B3LYP/6-31+G(d,p) with water (started 28 Oct), then add counter-ions, a longer chain and the glycolipid anchor.')
    rep(find_p(d,'The LTA fit in Section 6.4'),'The LTA fit in Section 6.4 was to April data, not to the new library reference.','The LTA fit in Section 6.4 was to April data; the comparison document tests the fragments and the model against the v2.1 reference. The LTA model (Section 6.7) is a quick-level, three-unit, gas-phase calculation.')
    rep(find_p(d,'For LPS or LTA fragments use the route'),'For the lipid A model use','For the LTA model use lta_model.gjf (two-job input) and make_spectrum.py (Section 6.6). For the lipid A model use')
    t=tb(d,'Content','Location'); addrow(t,['LTA whole-molecule model','C:\\Users\\sukes\\Downloads\\sukesh\\simulation\\lta output simulation (copy in SERS_Sepsis_Spectral_Library\\08_Methodology\\simulation\\LTA_Model_20261028): lta_model.gjf, .sdf, .smi, .out, lta_model_raman_act.csv, lta_model_spectrum_785nm.csv/.png, build_lta.py, make_spectrum.py, LTA_Raman_Methods.md, LTA_Raman_Full_Report.docx'])
    d.save(os.path.join(REP,F['sim']))
dft(); sim(); print('done')
