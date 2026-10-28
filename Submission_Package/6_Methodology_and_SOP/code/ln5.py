import os,pickle,shutil,numpy as np,pandas as pd,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from lib2 import LIB
N=pickle.load(open('ln1.pkl','rb'));Q=pickle.load(open('ln3.pkl','rb'));df=N['df'];P=N['P'];x=N['x']
GZ=(x>=1308)&(x<=1455);keep=~GZ
def nz(c): c=c-c[keep].min(); return c/c[keep].max()
TAG='NAM_Session_20260924'
# ---- tables
inv=df[['library_file','category','power','original_folder','original_file','md5','copy_verified','acq_time','note']].copy()
inv['original_folder']='bulk-lps-24-09-26/'+inv.original_folder
inv['identity_note']=np.where(inv.category=='NAM_Bulk','Folder named "lps" in error; operator checked bottle 27 Sep 2026: NAM; spectra match NAM reference','Empty glass slide blank')
qc=df[['library_file','category','power','integration_s','accumulations','dark_median','ripple_ds','raw_max_400_1800','n_sat_400_1800','sat_from','sat_to','noise','h929_counts','snr929','r_NAM_ref','QC','QC_reason']].round(3)
GSx=Q['SS'].round(4); BT=Q['BT'].round(3); CMP=Q['CMP'].round(3)
# bulk intensity table by condition (929 band + primary markers)
bi=[]
for name in ['REF_90_30May_all15','REF_90_30May_indep5','NEW_60_24Sep','NEW_70_24Sep','NEW_80_24Sep','NEW_90_24Sep']:
  for _,r in BT[(BT.set==name)&BT.role.str.startswith('Primary')].iterrows():
    bi.append(dict(set=name,band_ref=r.ref_pos,pos_mean=r.pos_mean,pos_sd=r.pos_sd,height_mean_counts=r.height_mean,height_cv_pct=r.height_cv,snr_mean=r.snr_mean,detected=r.detected))
bi=pd.DataFrame(bi)
dec=pd.DataFrame([
 ['Sample identity','Operator confirmed on 27 Sep 2026 that the powder measured on 24 Sep was NAM (folder named "lps" in error). Spectra match the NAM reference (r = 0.98–0.99 for set means).','Files renamed NAM_Bulk_20260924_*; the LPS screening outputs were withdrawn (moved to _to_delete/LPS_mislabel_20260924_superseded).'],
 ['New data vs existing reference','21 QC-passed independent spots (60 %: 11, 70 %: 5, 80 %: 4, 90 %: 1). All 7 primary and 4 supporting markers detected in 21/21; position shift ≤ 1.5 cm-1 (< 1 pixel); relative intensity (I/I929) within ±7 % for primary markers.','SUPPORTS the existing NAM reference (band positions, relative intensities, marker hierarchy).'],
 ['Reference condition','Absolute intensity differs between sessions (929 band at 90 %: 919 counts on 30 May vs 709 counts, n = 1, on 24 Sep) — not a reason to change the condition.','Reference condition unchanged: 90 % / 30 s / 5 acc.'],
 ['Top-up at 90 %','Only 1 of 13 new spectra at 90 % passed QC (dark-reference ripple).','Reference remains INTERIM (6 independent spots at 90 % incl. the new S09). Collect ≥ 4 more clean independent spots at 90/30/5 (target 10) with a fresh dark before every spectrum.'],
 ['Glass blanks','16 matched blanks, all saturated across the glass band.','Glass-zone bands (1345, 1392, 1450) stay "uncertain – inside glass zone"; add one labelled short-integration blank next session.'],
 ['LPS','No LPS data in this session.','LPS screening (70/80/90 % × 30 s × 5 acc + blanks) still required; the earlier "70 % provisional" LPS recommendation is withdrawn because it was based on NAM data.']],columns=['Item','Finding','Decision'])
# ---- write sheets into main workbook
wbp=os.path.join(LIB,'05_Tables','Spectral_Library_Metadata_and_Peaks.xlsx')
shutil.copy2(wbp,os.path.expanduser('~/work/wb_backup_before_20260924.xlsx'))
with pd.ExcelWriter(wbp,engine='openpyxl',mode='a',if_sheet_exists='replace') as w:
  for n,t in [('26_NAM_20260924_Inventory',inv),('27_NAM_20260924_QC',qc),('28_NAM_20260924_SetStats',GSx),('29_NAM_20260924_Old_vs_New',CMP),('30_NAM_20260924_Bands',BT),('31_NAM_Bulk_Intensity_Power',bi.round(2)),('32_NAM_20260924_Decision',dec)]:
    t.to_excel(w,sheet_name=n,index=False)
wb=openpyxl.load_workbook(wbp)
for n in wb.sheetnames[-7:]:
  ws=wb[n]
  for c in ws[1]: c.font=Font(bold=True);c.fill=PatternFill('solid',fgColor='D9D9D9');c.alignment=Alignment(wrap_text=True,vertical='top')
  for col in ws.columns:
    L=max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:60]);ws.column_dimensions[col[0].column_letter].width=min(max(10,L*0.9),70)
  ws.freeze_panes='B2'
ws=wb['README']; r=ws.max_row+1
while r>1 and ws.cell(r-1,1).value is None: r-=1
for t in ['','UPDATE 27 Sep 2026 — NAM session 24 Sep 2026 added (sheets 26–32). The raw folder "bulk-lps-24-09-26" contains NAM (operator-confirmed); files renamed NAM_Bulk_20260924_* / Glass_Blank_20260924_*.',
  'Result: 21/41 NAM spectra pass QC (failures = dark-reference ripple). New independent spots SUPPORT the NAM reference (positions ≤1.5 cm-1, relative intensities within ±7 % for primary markers). Reference condition unchanged (90/30/5); reference still INTERIM pending ≥4 more clean independent spots at 90 %.',
  'LPS: no LPS data yet — screening still required.']:
  ws.cell(r,1,t); r+=1
wb.save(wbp)
# ---- reference folder: validation CSVs + README
RD=os.path.join(LIB,'03_Reference_Spectra','NAM')
for pw in (60,70,80):
  L=Q['sets'][f'NEW_{pw}_24Sep'];C=np.array([P[f]['p']['corr'] for f in L]);Nm=np.array([nz(c) for c in C])
  pd.DataFrame({'Raman_Shift_cm-1':x,'BaselineCorrected_mean_counts':C.mean(0),'BaselineCorrected_median_counts':np.median(C,0),'BaselineCorrected_SD_counts':C.std(0,ddof=1),
    'Normalised_mean':Nm.mean(0),'Normalised_median':np.median(Nm,0),'Normalised_SD':Nm.std(0,ddof=1)}).round(5).to_csv(os.path.join(RD,f'NAM_Validation_20260924_P{pw:03d}pct_T30s_A05_Mean_Median_SD.csv'),index=False)
CMP.to_csv(os.path.join(RD,'NAM_Validation_20260924_Marker_Comparison.csv'),index=False)
with open(os.path.join(RD,'README_NAM_reference_status.txt'),'a') as fh:
  fh.write('''
UPDATE 27 Sep 2026 - independent-spot validation session 24 Sep 2026 (folder bulk-lps-24-09-26, operator-confirmed NAM)
- 41 NAM spectra at 60/70/80/90 % x 30 s x 5 acc + 16 matched glass blanks; 21 pass QC (60 %: 11/11, 70 %: 5/8, 80 %: 4/9, 90 %: 1/13).
  All failures = dark-reference ripple (dark level rose ~1420 -> ~4470 counts during the session); no sample saturation.
- Validation: set means correlate with the reference mean at r = 0.982-0.994 (glass zone excluded). All primary and supporting markers
  detected in 21/21 spectra; position shift <= 1.5 cm-1 (< 1 detector pixel); relative intensity I/I929 within +/-7 % (primary),
  +/-11 % (other bands outside glass zone); glass-zone bands +1 to +13 % (variable glass contribution).
- Conclusion: the new data SUPPORT the existing reference. Reference condition unchanged (90 % / 30 s / 5 acc). Reference files unchanged.
- Status: still INTERIM - independent spots at 90/30/5 = 5 (30 May) + 1 (24 Sep, S09). Collect >= 4 more clean independent spots
  (target 10) at 90/30/5 with a fresh dark before every spectrum.
- Files: NAM_Validation_20260924_*.csv (this folder); workbook sheets 26-32; report NAM_Bulk_Validation_Report_20260924.docx.
''')
# ---- QC logs
qc.to_csv(os.path.join(LIB,'07_QC_and_Exclusions',f'{TAG}_QC_Log.csv'),index=False)
ex=df[(df.category=='NAM_Bulk')&(df.QC!='PASS')]
rows=pd.DataFrame({'File name':ex.library_file,'Original file':'bulk-lps-24-09-26/'+ex.original_folder+'/'+ex.original_file,'Class':'NAM','Power (%)':ex.power,'Integration (s)':ex.integration_s,'Acc':ex.accumulations,
  'Decision':'Excluded (QC)','Reason':ex.QC_reason,'Library folder':'NAM_Bulk'})
rows.to_csv(os.path.join(LIB,'07_QC_and_Exclusions','Excluded_and_Flagged_Spectra.csv'),mode='a',header=False,index=False)
print('ok',len(rows))
