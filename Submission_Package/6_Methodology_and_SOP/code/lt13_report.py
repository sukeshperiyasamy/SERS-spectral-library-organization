# Builds 09_Reports/LTA_Bulk_Reference_Report_20261028.docx (version 2.0, all data of the 28 Oct 2026 session)
import os,pickle,numpy as np,pandas as pd
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lib2 import LIB,X
H=os.path.dirname(os.path.abspath(__file__)); R=pickle.load(open(os.path.join(H,'lt11_result.pkl'),'rb')); DK=pd.read_pickle(os.path.join(H,'lt10_dark.pkl'))
d,rec,PL,bands,ct=R['d'],R['rec'],R['PL'],R['bands'],R['ct']; POOL=set(R['POOL']); cmp_=R['cmp']; FIG=os.path.join(LIB,'04_Figures','LTA_Reference_20261028')
GZ=(X>=1308)&(X<=1455)
sp=d[d.role=='spot']; ss=d[d.role=='samespot']; bl=d[(d.role=='blank')&(d.QC!='BLANK-SUSPECT')]; mem=sp[sp.lib.isin(POOL)]; alll=d[d.role!='blank']
doc=Document(); s=doc.sections[0]; s.page_width=Cm(21); s.page_height=Cm(29.7)
for k in ('left_margin','right_margin','top_margin','bottom_margin'): setattr(s,k,Cm(1.4))
st=doc.styles['Normal']; st.font.name='Arial'; st.font.size=Pt(10); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Arial'); st.paragraph_format.space_after=Pt(4)
for n,sz in (('Heading 1',14),('Heading 2',11.5)):
    h=doc.styles[n]; h.font.name='Arial'; h.font.size=Pt(sz); h.font.color.rgb=RGBColor(0,0,0); h.element.rPr.rFonts.set(qn('w:eastAsia'),'Arial')
def shade(c,hx):
    pr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),hx); pr.append(sh)
def para(t,bold=False,italic=False,size=None):
    p=doc.add_paragraph(); r=p.add_run(t); r.bold=bold; r.italic=italic
    if size: r.font.size=Pt(size)
def bullets(items):
    for t in items: doc.add_paragraph(t,style='List Bullet')
def table(head,rows,widths,fs=8.5,fills=None):
    t=doc.add_table(rows=1,cols=len(head)); t.style='Table Grid'
    for i,x in enumerate(head):
        c=t.rows[0].cells[i]; c.text=''; r=c.paragraphs[0].add_run(x); r.bold=True; r.font.size=Pt(fs); shade(c,'D9D9D9')
    for k,row in enumerate(rows):
        cs=t.add_row().cells
        for i,x in enumerate(row):
            cs[i].text=''; r=cs[i].paragraphs[0].add_run(str(x)); r.font.size=Pt(fs)
        if fills and fills[k]: shade(cs[-1],fills[k])
    for row in t.rows:
        for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    doc.add_paragraph()
FN=[0]
def fig(n,cap,w=17):
    FN[0]+=1; doc.add_picture(os.path.join(FIG,n),width=Cm(w)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER; para(f'Figure {FN[0]}. '+cap,italic=True,size=8.5); return FN[0]
cond=lambda r:f"{int(r.power)} %, {r.int_s:g} s, {int(r.acc)} acc"
nleak=int((alll.QC=='FLAG-DARK').sum()); nl=len(alll)
c80_30=ct[(ct.power==80)&(ct.int_s==30)].iloc[0]; c80_45=ct[(ct.power==80)&(ct.int_s==45)].iloc[0]; c70_60=ct[(ct.power==70)&(ct.int_s==60)].iloc[0]
sp30=sp[(sp.power==80)&(sp.int_s==30)&(~sp.leak)]
# ---------------- title
para('LTA Bulk-Powder Raman Reference: Full Session of 28 October 2026 on Glass',bold=True,size=17)
para('Quality control, reference spectrum, bands, parameters tested, and what to change',size=11)
para(f'SERS sepsis project, bulk-powder Raman spectral library. Version 2.0, 28 October 2026. Raw data: bulk-lta ({len(alll)} LTA spectra and {len(bl)} glass blanks used). Supersedes version 1.0 of this report; the preliminary April 2026 reference (4 spectra) is kept unchanged.',size=9)
doc.add_heading('1. Summary',1)
para(f"The session gave {len(alll)} LTA spectra and {len(bl)} blanks at nine laser, time and accumulation settings. {len(mem)} independent spots pass all quality checks and form the new reference (v2.1): 10 at 80 % power, 45 s and 2 accumulations, and 7 at 80 % power, 30 s and 3 accumulations. Its mean spectrum matches the earlier 80 %, 45 s reference (r = {cmp_['r_pool_vs_v2']:.3f}), the April data, and is clearly different from NAM (r = {cmp_['r_pool_vs_NAM']:.2f}) and LPS (r = {cmp_['r_pool_vs_LPS']:.2f}). The strongest band, 1092 cm-1, is present on every reference spot at a signal-to-noise ratio of about 30.")
para(f"Three findings matter for the next session. First, 80 % power, 30 s and 3 accumulations is the better setting: all 7 spots measured with a clean dark passed, none saturated, against 10 of 28 at 45 s and 2 accumulations. Second, {nleak} of the {nl} LTA spectra ({100*nleak/nl:.0f} %) carry a dark frame that contains sample light, because the dark was not retaken after the settings changed; they are flagged and kept out of the reference, and every low-power result of the day is among them. Third, the blanks at 80 %, 30 s and 3 accumulations are clean glass; the blanks at 45 s and 60 s saturate.")
table(['Item','Result'],[
 ['Condition of the reference','785 nm, 80 % power; 10 spots at 45 s x 2 acc and 7 spots at 30 s x 3 acc; shutter-closed dark'],
 ['Reference v2.1 (pooled)',f"{len(mem)} spots; mean and median agree (r = {cmp_['mm']:.3f}); mean pairwise r {cmp_['pair_mean']:.2f}; the two settings agree (r = {cmp_['r_30_vs_45']:.3f})"],
 ['Reference v2 (45 s only)','10 spots, kept unchanged; pooled mean agrees with it at r = %.3f'%cmp_['r_pool_vs_v2']],
 ['Primary bands','1092 cm-1 (SNR 30), 539 cm-1 (SNR 17), 862 cm-1 (SNR 16)'],
 ['Supporting bands','419, 723, 798, 1010, 1459 and 1751 cm-1'],
 ['Best setting tested',f"80 %, 30 s, 3 acc: {int(c80_30.PASS)} of {int(sp30.shape[0])} clean-dark spots pass; no saturation; ripple 29 to 46 counts"],
 ['Dark frame problem',f"{nleak} of {nl} LTA spectra flagged (dark median 2,500 to 11,500 counts against a clean 1,333 to 1,680)"],
 ['Sample','Sigma-Aldrich L2515-10MG, LTA from Staphylococcus aureus, lot 0000419622, source number 0000404178; same vial as April 2026']],[4.2,13.8])
doc.add_heading('2. Sample and instrument',1)
table(['Item','Detail'],[['Sample','Lipoteichoic acid from Staphylococcus aureus, Sigma-Aldrich L2515-10MG, lot 0000419622, source number 0000404178 (as given by the operator; the label reads bacterial cell wall polymer); powder placed directly on a glass slide. The same vial was used for the 3 Apr 2026 data. The supplier product sheet gives no purity figure and no physical form; the certificate of analysis for the lot has not been obtained, so purity is Not recorded.'],
 ['Supplier description','Linear polymer of phosphodiester-linked glycerol phosphate bound to a lipid; the C2 position usually glycosylated and/or D-alanylated; soluble in water at 5 mg/ml; stored at 2 to 8 C (product information sheet, summarised in LTA_Sample_Product_Information_L2515.txt)'],
 ['Instrument','B&W Tek i-Raman Plus, BWS465-785H, 785 nm; usable range ends at 2842 cm-1; objective not recorded'],
 ['Laser power','Instrument setting in % of the maximum (not converted to mW)'],
 ['Date and time','28 Oct 2026, 02:30 to 11:02 (instrument clock)']],[4,14])
doc.add_heading('3. What was measured',1)
para('All files come from the folder bulk-lta. The instrument numbered the files in two runs (SP_1 to SP_36, then SP_1 to SP_82), so the original names repeat; the library names add the date, power, time and accumulations to make each one unique. In the library, spots sit in one folder per setting under 01_Raw_Data_Renamed/LTA_Bulk, and the blanks in Glass_Blank/20261028_LTA_session.')
grp=[]
for sub,g in d[d.QC!='BLANK-SUSPECT'].groupby('sub_lib'):
    a=g[g.role=='spot']; b=g[g.role=='samespot']; c=g[(g.role=='blank')&(g.QC!='BLANK-SUSPECT')]
    conds=sorted(set(cond(r) for r in g[g.role!='blank'].itertuples())); t0=g.acq.min()[11:16]; t1=g.acq.max()[11:16]
    grp.append([sub.replace('2026-10-28_','').replace('_',' '),len(a),len(b),len(c),'; '.join(conds),f'{t0} to {t1}'])
table(['Library folder (2026-10-28_...)','Independent spots','Same-spot scans','Blanks','Settings','Time'],grp,[4.6,2,2,1.6,5,2.8],fs=8)
table(['Original folder in bulk-lta','Spectra','What it is'],[['45sec-80p-2acc-diffrentspot','28 used','28 spots at 80 %, 45 s, 2 acc (SP_9-36)'],['diifrent 60s-2ac-70p','9','1 spot at 50 %, 40 s; 8 spots at 70 %, 60 s, 2 acc'],['blank glass','11','Blanks: 70 %/60 s (2), 80 %/45 s (5), 80 %/30 s x3 (4)'],['30sec-80p-3ac','20','7 spots and 13 same-spot scans at 80 %; 7 scans at 30 s x 3 and 6 scans at 45 s x 2'],['70power-60power mulispot','20','Power series at 45 s x 2: 70, 60, 50, 40 and 30 %'],['30s-80p-3acc','11','5 spots and 6 same-spot scans at 80 %, 30 s, 3 acc'],['30sec-70p-3 accc','11','5 spots and 6 same-spot scans at 70 %, 30 s, 3 acc']],[5,2.6,10.4],fs=8)
doc.add_heading('4. Parameters and processing',1)
table(['Parameter','Setting'],[['Settings tested','80 %: 45 s x 2 and 30 s x 3; 70 %: 60 s x 2, 45 s x 2 and 30 s x 3; 60, 50, 40 and 30 % at 45 s x 2; one spot at 50 %, 40 s x 2'],
 ['Processing','Pipeline v3: dark-subtracted signal; two defective detector pixels repaired; 400-1800 cm-1; despiking (MAD z > 6, width 2 pixels or less); Savitzky-Golay 11/3; ALS baseline (lambda 1e5, p 0.01, 10 iterations); noise = 1.4826 x MAD; min-max normalisation'],
 ['QC limits','No saturation in 400-1800 cm-1; ripple 50 counts or less; SNR of 1092 at least 10; glass-dominated spots excluded (area ratio 8 or more); atypical shape excluded (r to the reference mean below 0.80)'],
 ['Dark check','Dark median must be near the clean value for the integration time: 1,333 (30 s x 3), 1,443 (40 s x 2), 1,510 (45 s x 2), 1,680 (60 s x 2) counts; more than 150 counts above = dark leak, flagged and not used in the reference'],
 ['Same-spot scans','Repeat scans on one spot are not independent and are never part of a reference'],
 ['Band classes','C1 primary (all spots, SNR 15 or more, glass share below 25 %, ratio CV 30 % or less); C2 supporting (80 % of spots, SNR 5 or more); C3 glass; C4 uncertain'],
 ['Glass zone','1308-1455 cm-1, left out of all correlations']],[3.6,14.4])
doc.add_heading('5. Findings',1)
doc.add_heading('5.1 Quality control by setting',2)
rows=[[f"{int(r.power)} %, {r.int_s:g} s, {int(r.acc)} acc",int(r.n),int(r.PASS),int(r.REVIEW),int(r.EXCLUDE),int(r.FAIL),int(r.FLAG_DARK),f"{r.dark_med:.0f}",f"{r.ripple_med:.0f}",f"{r.snr_med:.0f}"] for r in ct.sort_values(['power','int_s'],ascending=[False,True]).itertuples()]
table(['Setting','Spots','PASS','REVIEW','EXCLUDE','FAIL','Dark leak','Dark median','Ripple (median)','SNR 1092 (median)'],rows,[3.4,1.2,1.2,1.5,1.7,1.2,1.6,1.8,1.8,2.2],fs=8)
fig('FigL5_QC_Overview.png','QC outcome of the independent spots at every setting (left) and the pass rate of the spots with a clean dark (right). Dark leak = the dark frame contains sample light.')
para(f"With a clean dark, 80 %, 30 s and 3 accumulations is clearly the best: {int(sp30.shape[0])} of {int(sp30.shape[0])} spots pass, the highest raw count is {max(rec[x]['rawmax'] for x in sp30.rel):,.0f} (no saturation), ripple is {min(rec[x]['ripple'] for x in sp30.rel):.0f} to {max(rec[x]['ripple'] for x in sp30.rel):.0f} counts, and the SNR of the 1092 band is {min(rec[x]['snr'] for x in sp30.rel):.0f} to {max(rec[x]['snr'] for x in sp30.rel):.0f}. At 80 %, 45 s and 2 accumulations only 10 of 28 spots pass: 5 saturate, 2 are weak, 3 have an atypical shape and 8 have ripple of 50 to 58 counts. At 70 %, 60 s and 2 accumulations (clean dark) 2 of 8 pass and 5 have ripple of 51 to 61 counts, so lowering the power and lengthening the time does not help. The other settings cannot be judged because their dark frames are contaminated (Section 5.2).")
fig('FigL9_Condition_Comparison.png','Highest raw count, ripple, SNR of the 1092 band and its height for the settings with a clean dark.',w=16.5)
doc.add_heading('5.2 The dark frame problem',2)
dd=[]
for (pw,it,ac),g in alll.groupby(['power','int_s','acc']):
    dd.append([f"{int(pw)} %, {it:g} s, {int(ac)} acc",len(g),f"{np.median([rec[x]['dark'] for x in g.rel]):.0f}",f"{CLEANV.get((it,ac),'')}" if False else ''])
CLEANV={(30.0,3):1333,(40.0,2):1443,(45.0,2):1510,(60.0,2):1680}
dd=[[f"{int(pw)} %, {it:g} s, {int(ac)} acc",len(g),f"{min(rec[x]['dark'] for x in g.rel):.0f} to {max(rec[x]['dark'] for x in g.rel):.0f}",CLEANV.get((it,ac),''),int(g.leak.sum())] for (pw,it,ac),g in alll.groupby(['power','int_s','acc'])]
table(['Setting','Spectra','Dark median in the files (counts)','Clean dark (counts)','Spectra with a dark leak'],dd,[3.6,1.8,4.6,3,3],fs=8.5)
para(f"The dark frame is the reference level the instrument subtracts from every spectrum. For {nl-nleak} spectra it is clean, at the level expected for the integration time. For {nleak} spectra it is far higher (2,500 to 11,500 counts), which means the dark was recorded with sample light on the detector, or an old dark from another setting was kept. The first spectra of the day (02:30 to 07:30) had clean darks; from the same-spot tests at 07:35 onward the dark level climbs in steps through the rest of the day (Figure 3).")
bullets(['Subtracting a dark that contains sample light removes part of the real signal, so the intensities of these spectra are too low (the 1092 band comes out at about half its height in the test below) and the ripple measure is unreliable.',
 'A test with a clean dark of the same integration time put back in did not repair them: the median ripple rose (for example 68 to 212 counts at 80 %, 45 s) and the shape agreement with the reference fell (r about 1.0 to 0.8). The ripple fringes are in the collected light, so the contaminated dark was cancelling them. The flagged spectra are therefore kept as a record and not corrected.',
 'Every spectrum at 70 % and below, except the 70 %, 60 s spots, falls in this group. The power-series results (Section 5.5) cannot be used for conclusions, and need repeating with a fresh dark at each setting.'])
fig('FigL10_Dark_Check.png','Dark level of every LTA spectrum through the day (left; dotted lines are the clean levels) and the SNR of the 1092 band against the dark excess at 45 s x 2 (right).')
doc.add_heading('5.3 The blanks',2)
para('Eleven glass blanks are used (05:29 to 05:59). They match the 28 September glass blanks (r = 0.69 to 0.98). However, those at 80 %, 45 s and at 70 %, 60 s saturate the glass band (65,535 counts), so they cannot serve as full-condition blanks. The 4 blanks at 80 %, 30 s and 3 accumulations are clean (r = 0.98, highest raw count about 60,000, no saturation); their mean was used as the glass shape in the glass-share estimate.')
fig('FigL13_Blanks.png','The 11 glass blanks. Orange: saturated. Green: clean glass at 80 %, 30 s, 3 acc.')
doc.add_heading('5.4 Same-spot repeat scans',2)
para('Four series of repeat scans on one spot were taken (6 or 7 scans each). All four have a leaked dark, so intensities are low and ripple is unreliable, but the shape can still be compared. At 80 %, 30 s and 3 accumulations the height of 1092 stays within about 8 % of the first scan over 6 to 7 scans and the 1459/1092 ratio does not change, which gives no sign of damage to the powder at this power. The series at 80 %, 45 s and 2 accumulations drifts (1092 up by 23 %, ratio down) in step with a rising dark level and ripple, so that change cannot be assigned to the sample. These repeat scans are not independent spots and are kept out of the reference.')
fig('FigL11_Same_Spot_Stability.png','Same-spot series: height of 1092 relative to the first scan (left) and the 1459/1092 ratio (right). All series have a leaked dark.',w=16.5)
doc.add_heading('5.5 Power series at 45 s',2)
para('The multi-spot series at 70, 60, 50, 40 and 30 % (45 s, 2 acc; 20 spots) all have leaked darks (2,900 to 10,000 counts). The 1092 band is found at every power, with a signal-to-noise ratio of 9 to 23, but the intensities and ripple cannot be compared with the clean-dark spots. No conclusion is drawn about how the signal scales with power. The series should be repeated with a fresh shutter-closed dark at each power.')
fig('FigL12_Power_Series.png','Height of 1092 (left) and highest raw count (right) against laser power at 45 s x 2. Purple: dark with sample light.',w=16.5)
doc.add_heading('5.6 Reference spectrum v2.1',2)
para(f"The reference is the mean of the {len(mem)} passing spots of the two clean-dark settings. The mean agrees with the median (r = {cmp_['mm']:.3f}), the 10 spots at 45 s and the 7 at 30 s give mean spectra that agree at r = {cmp_['r_30_vs_45']:.3f}, mean pairwise correlation between spots is {cmp_['pair_mean']:.2f} (lowest {cmp_['pair_min']:.2f}), and the pooled mean agrees with the 10-spot v2 reference at r = {cmp_['r_pool_vs_v2']:.3f}. The normalised shape is therefore the same at both settings; absolute counts differ and are compared as band ratios.")
fig('FigL3_LTA_Reference_Bands.png','Reference v2.1: mean, median and SD of 17 spots with the classified bands. The dashed line is the clean glass blank scaled to the median spot.')
mrows=[[r.lib.split('_')[-1].replace('.csv','').replace('S',''),f"{r.power} %, {r.int_s:g} s, {r.acc} acc",r.acq[11:16],f"{rec[r.rel]['rawmax']:.0f}",f"{rec[r.rel]['ripple']:.0f}",f"{rec[r.rel]['snr']:.0f}",f"{rec[r.rel]['rREF']:.2f}"] for r in mem.itertuples()]
table(['Spot (S)','Setting','Time','Highest raw count','Ripple','SNR 1092','r to v2 mean'],mrows,[1.6,3.4,1.8,3,1.8,2,2.4],fs=8)
doc.add_heading('5.7 Bands and comparison with earlier data',2)
rows=[]
for r in bands.itertuples():
    c=r._1[:2]
    if c in('C1','C2') or r.band in(629,1247,1583): rows.append([f"{r.pos_mean:.1f}",{'C1':'Primary','C2':'Supporting','C4':'Uncertain'}.get(c,c),r.detected,f"{r.pos_sd:.1f}",f"{r.snr_median:.0f}",f"{r.height_cv:.0f} %",f"{r.glass_share_pct:.1f} %"])
table(['Band (cm-1)','Class','Found in','Position SD','SNR (median)','Height CV','Glass share (estimate)'],rows,[2,2.4,2,2,2.4,2.2,3.4],fs=8.5)
para('Three bands are primary (1092, 539, 862) and six supporting (419, 723, 798, 1010, 1459, 1751). 1247 cm-1 (phosphate region) and 629 and 1583 cm-1 are seen on 12, 13 and 11 of 17 spots and stay in the uncertain class. The glass share is estimated from the clean glass blank scaled to each spot on 1290 to 1520 cm-1; it is below 7 % at every band, so the slide contributes little under the powder. Against the April reference the bands agree within 3 cm-1 (802 to 799, 861.5 to 862, 1091.8 to 1092, 1458.8 to 1460), and 541 (uncertain in April) is now a reproducible primary band at 539.')
fig('FigL4_LTA_vs_Other_References.png','Pooled LTA reference against the April LTA, LPS and NAM references (normalised).')
# ---------------- peaks section
doc.add_heading('6. What the peaks say and how they are related',1)
para('Each labelled peak is read against the structure the supplier gives for S. aureus LTA: a linear polymer of phosphodiester-linked glycerol phosphate bound to a lipid, with the C2 position usually glycosylated and/or D-alanylated. All assignments are tentative: they say what a band is consistent with, not what it proves.')
fig('FigL6_LTA_Peaks_Annotated.png','Reference v2.1 with every labelled peak, coloured by the part of the molecule it is consistent with: red phosphodiester, blue glycerol skeleton, orange lipid chain, purple ester C=O, grey not assigned.')
def hh(rel,p,t=6): w=(X>=p-t)&(X<=p+t); return float(rec[rel]['corr'][w].max())
PK=[419,539,629,723,798,862,1010,1092,1247,1459,1583,1751]; M=np.array([[hh(r.rel,p) for p in PK] for r in mem.itertuples()]); idx={p:i for i,p in enumerate(PK)}
RAT={p:(float(np.median(M[:,idx[p]]/M[:,idx[1092]])),100*float(np.std(M[:,idx[p]]/M[:,idx[1092]],ddof=1)/np.mean(M[:,idx[p]]/M[:,idx[1092]]))) for p in PK}
CR=np.corrcoef(M.T); COR=lambda a,b:float(CR[idx[a],idx[b]])
def rt(p): return f"{RAT[p][0]:.2f} (CV {RAT[p][1]:.0f} %)"
doc.add_heading('6.1 Peak by peak',2)
table(['Band (cm-1)','Found; height relative to 1092','Part of LTA it is consistent with','What the peak says','Caution'],[
 ['1092','17/17; 1.00','PO2- symmetric stretch of the phosphodiester, one in every repeat','The phosphate backbone is present. It is the strongest and most reliable band, and its height follows the amount of backbone in the focus','Phosphate groups also occur in phospholipids and nucleic acids, so the band alone is not specific (group frequency 1090-1095: Movasaghi 2007; Czamara 2015; Grunert 2018)'],
 ['1247','12/17; '+rt(1247),'PO2- asymmetric stretch region (1225-1250)','The second band of the phosphate pair; with 1092 it completes the PO2- signature','Weaker and broader; a protein impurity would add a band in the same region (purity not recorded)'],
 ['798','16/17; '+rt(798),'O-P-O / C-O-P stretching of the link between glycerol units','Phosphate bridges between the glycerol units','Region only; the April literature value of 813 cm-1 is 14 cm-1 higher; glass contributes under 1 %'],
 ['862','17/17; '+rt(862),'C-C and C-O-C stretching of the glycerol (and any sugar) skeleton','The glycerol skeleton of the repeat','Carbohydrate region 800-950 (Wiercigroch 2017); not specific'],
 ['539','17/17; '+rt(539),'Skeletal deformation of the backbone (C-C-O, O-P-O)','Low-frequency fingerprint of the whole backbone; reproducible on every spot','No checked literature band; region only'],
 ['419','17/17; '+rt(419),'Skeletal deformation','Low-frequency fingerprint','Region only; close to the 425 cm-1 band of LPS, so use it only with the rest of the pattern'],
 ['723','17/17; '+rt(723),'CH2 rocking of long alkyl chains (general group frequency, not in the library tables)','The lipid anchor chains','Tentative; follows 1459'],
 ['1010','14/17; '+rt(1010),'C-O / C-C stretching region','Not interpreted further','New in this session; no checked literature band; the clean glass blank has no band here'],
 ['1459','16/17; '+rt(1459),'CH2 scissoring of the lipid chains and glycerol CH2','Hydrocarbon is present (lipid anchor, glycerol)','Non-specific (Czamara 2015)'],
 ['1751','16/17; '+rt(1751),'Ester C=O stretching region','Ester links, such as D-alanyl esters on the glycerol and/or fatty-acid esters of the lipid','Weak; ester region from lipid group frequencies (Czamara 2015)'],
 ['629, 1583','13/17 and 11/17','Not assigned','No interpretation','Not seen in the clean glass blank; check again with a fresh clean blank'],
 ['1328','17/17','Glass zone','Glass, not LTA','Never used as an LTA band']],[1.7,2.4,4.3,5,4.6],fs=7.5)
doc.add_heading('6.2 How the peaks are related',2)
bullets(["By chemical group. The two phosphate bands (1092 and 1247) belong together as the symmetric and asymmetric stretch of the same PO2- group; 798 adds the O-P-O link. 862, 539 and 419 describe the glycerol skeleton. 1459 and 723 come from CH2 in the lipid chains, and 1751 from the ester links. The structure map in Figure %d places these groups on the molecule."%(FN[0]+2),
 f"Between spots. The heights of all bands rise and fall together (correlation with 1092: 798 r = {COR(1092,798):.2f}, 862 r = {COR(1092,862):.2f}, 1459 r = {COR(1092,1459):.2f}, 1751 r = {COR(1092,1751):.2f}, 1247 r = {COR(1092,1247):.2f}, 539 r = {COR(1092,539):.2f}), because every band depends on how much powder is in the focus. This alone does not show a common origin.",
 f"Ratios to 1092. More informative is that the ratios stay steady from spot to spot: 798 and 862 vary least (CV {RAT[798][1]:.1f} % and {RAT[862][1]:.1f} %), the lipid bands 1459 and 723 a little more ({RAT[1459][1]:.0f} % and {RAT[723][1]:.0f} %), and the weakest bands most (1247 about {RAT[1247][1]:.0f} %). A steady ratio is what one material gives, and the tight ratios of the backbone bands fit their sharing one unit.",
 f"Backbone to lipid. 1459 follows 1092 closely (r = {COR(1092,1459):.2f}) and 723 follows 1459 (r = {COR(1459,723):.2f}); the lipid-to-backbone ratio (1459/1092 = {RAT[1459][0]:.2f}) is steady, as expected when the lipid is covalently bound to the polymer and not a separate component.",
 f"Ester to chain. 1751 follows both 1459 and 1092 (r = {COR(1459,1751):.2f} and {COR(1092,1751):.2f}), in line with ester links on the polymer."])
table(['Band (cm-1)','Median height / height at 1092','CV between spots'],[[p,f"{RAT[p][0]:.2f}",f"{RAT[p][1]:.1f} %"] for p in (798,862,539,1459,723,419,1247,1751)],[3,5,4],fs=8.5)
fig('FigL7_LTA_Peak_Zooms.png','The main bands in the 17 reference spots (thin lines) and their mean (thick line).')
fig('FigL8_LTA_Structure_Map.png','Schematic of S. aureus LTA as the supplier describes it, with the bands that are consistent with each part. The drawing is generic and does not show chain length or which substituent sits on which glycerol.')
doc.add_heading('6.3 What the pattern says as a whole',2)
bullets(['A phosphodiester-linked glycerol backbone (1092, 1247, 798, 862), a hydrocarbon lipid part (1459, 723) and ester groups (1751) are all present in one reproducible spectrum. This is what the supplier describes for S. aureus LTA.',
 'It does not show which substituent (D-alanine or a sugar) is present, how long the chain is, or how pure the sample is. Protein or nucleic-acid impurities, and other phosphate lipids, would add bands in the same regions.',
 'It differs from LPS: the strongest LTA band is 1092 cm-1, whereas the strongest LPS band is 425 cm-1, and the overall correlation is %.2f. Both carry phosphate and sugar-like bands, which is why they are not unrelated. NAM has no phosphate and correlates at %.2f.'%(cmp_['r_pool_vs_LPS'],cmp_['r_pool_vs_NAM']),
 'Identity therefore rests on the supplier record (Sigma L2515, S. aureus, lot 0000419622) together with this pattern. The wording stays "consistent with LTA".'])
# ---------------- recommendations
doc.add_heading('7. What to change in the next session',1)
table(['Issue','Evidence','Change'],[
 ['Dark frame contains sample light',f'{nleak} of {nl} spectra; dark rose from 1,333 to 11,500 counts through the day','Take a new dark with the shutter closed by hand every time the power, time or accumulations change, and at least every 10 spots. Check the dark median in the file header before measuring: about 1,333 (30 s x 3), 1,443 (40 s x 2), 1,510 (45 s x 2), 1,680 (60 s x 2); retake if it is more than 150 counts higher'],
 ['Best setting','80 %, 30 s, 3 acc: 7 of 7 clean-dark spots pass, no saturation. 80 %, 45 s, 2 acc: 10 of 28 pass','Use 80 %, 30 s, 3 acc for new LTA spots, with a clean dark, and keep the reference at v2.1'],
 ['70 %, 60 s','2 of 8 pass; 5 with ripple 51 to 61','No gain over 80 %, 30 s; do not use'],
 ['Low-power series (30 to 70 %)','All 45 s and 30 s runs at 70 % and below carry a leaked dark','Repeat the series at 30 s x 3 with a fresh dark at each power if the power dependence is needed'],
 ['Blanks','Blanks at 45 s and 60 s saturate the glass band','Use blanks at 80 %, 30 s, 3 acc on the clean glass slide (clean, r = 0.98 to the 28 Sep glass); use a clean slide and avoid plastic tools'],
 ['Saturation at 45 s','5 of 28 spots saturate at 80 %, 45 s','Not a problem at 30 s x 3 (highest raw count 55,323); still skip spots whose short test shot is very bright'],
 ['Same-spot tests','All four series have a leaked dark','Repeat one series (6 scans) at 80 %, 30 s, 3 acc with a clean dark if damage must be excluded'],
 ['Sample record','Purity not recorded','Download the certificate of analysis for lot 0000419622 and record the purity'],
 ['Objective','Not recorded in any file','Write the objective and focus method on the recording sheet']],[3.5,5.5,9],fs=8)
doc.add_heading('8. Status and open points',1)
bullets([f'Status: ESTABLISHED, pending your confirmation. Reference v2.1 uses {len(mem)} spots from one session and one vial (the April data are the same vial).',
 'The main library report should be updated to v3.4 with a new section for this LTA reference once you confirm.',
 'The 55 flagged spectra stay in the library as a record. They can be replaced by new clean-dark spectra and need not be reprocessed.',
 'A second lot or session would test whether the bands are reproducible beyond this powder.'])
doc.add_heading('9. Files',1)
table(['Content','Location (SERS_Sepsis_Spectral_Library)'],[['Copies of the raw files','01_Raw_Data_Renamed/LTA_Bulk/2026-10-28_* (one folder per setting) and 2026-04-03_April_preliminary; Glass_Blank/20261028_LTA_session'],['Processed spectra','02_Processed_Spectra/LTA/ (same folders) and Glass_Blank/20261028_LTA_session'],['Reference v2 (10 spots) and v2.1 pooled (17 spots)','03_Reference_Spectra/LTA/ (files with 20261028 in the name; April files unchanged); status in README_LTA_reference_status.txt'],['QC log and exclusions','07_QC_and_Exclusions/LTA_Reference_Session_20261028_QC_Log.csv ; Excluded_and_Flagged_Spectra.csv'],['Figures','04_Figures/LTA_Reference_20261028 (FigL1 to FigL13)'],['Workbook sheets','05_Tables/Spectral_Library_Metadata_and_Peaks.xlsx, sheets 43 to 48; 05_Tables/LTA_20261028_condition_comparison.csv'],['Code','08_Methodology/code/lt1 to lt13'],['Original files','bulk-lta (unchanged)']],[5,13])
doc.save(os.path.join(LIB,'09_Reports','LTA_Bulk_Reference_Report_20261028.docx')); print('saved',FN[0],'figures')
