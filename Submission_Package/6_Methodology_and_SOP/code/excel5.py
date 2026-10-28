import sys,os,pickle,numpy as np,pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
sys.path.insert(0,os.path.expanduser('~/work')); from lib2 import *
FV=pickle.load(open(os.path.expanduser('~/work/final_v3.pkl'),'rb')); D=pickle.load(open(os.path.expanduser('~/work/an4.pkl'),'rb')); CT=pickle.load(open(os.path.expanduser('~/work/crit.pkl'),'rb')); SE=pickle.load(open(os.path.expanduser('~/work/sens.pkl'),'rb'))
fn=os.path.join(LIB,'05_Tables','Spectral_Library_Metadata_and_Peaks.xlsx'); wb=load_workbook(fn)
H=Font(name='Arial',bold=True,color='FFFFFF',size=10); HF=PatternFill('solid',fgColor='1F4E79'); F=Font(name='Arial',size=10); thin=Side(style='thin',color='BFBFBF')
def sheet(name,df,title,notes=(),widths=None,wrap=False):
  if name in wb.sheetnames: del wb[name]
  ws=wb.create_sheet(name); ws.cell(1,1,title).font=Font(name='Arial',bold=True,size=12); r0=2
  for n in notes: ws.cell(r0,1,n).font=Font(name='Arial',italic=True,size=9,color='52514E'); r0+=1
  r0+=1
  for j,c in enumerate(df.columns,1): x=ws.cell(r0,j,str(c)); x.font=H; x.fill=HF; x.alignment=Alignment(wrap_text=True,vertical='center')
  for i,row in enumerate(df.itertuples(index=False),r0+1):
    for j,v in enumerate(row,1):
      if isinstance(v,np.generic): v=v.item()
      if isinstance(v,bool): v='Yes' if v else 'No'
      if isinstance(v,float) and np.isnan(v): v=None
      c=ws.cell(i,j,v); c.font=F; c.border=Border(bottom=thin); c.alignment=Alignment(wrap_text=wrap,vertical='top')
  ws.freeze_panes=ws.cell(r0+1,2)
  for j,c in enumerate(df.columns,1):
    w=(widths or {}).get(c) or min(50,max(9,len(str(c))*0.9+2,*(len(str(v))+2 for v in df[c].head(300))))
    ws.column_dimensions[get_column_letter(j)].width=w
  return ws
sheet('Traceability_Table',FV['TR'],'CENTRAL TABLE — marker traceability: valid spectrum → reproducible → peak reproducible → not background → chemically plausible → literature',
 ['Every final marker (primary and supporting) is traced from QC-passed spectra to its literature comparison. Assignments are TENTATIVE; no pure-compound normal-Raman reference exists for NAM, LTA or LPS.','LPS: no marker (reference pending controlled glass-slide measurement).'],widths={'Assignment':55,'Evidence_type':45,'Replicate_frequency':38},wrap=True)
sheet('15_GO_NOGO',FV['GO'],'Sheet 15 — Formal GO / RE-ACQUIRE decision per biomarker (whole replicate group)',widths={c:40 for c in FV['GO'].columns[1:]},wrap=True)
sheet('16_Acquisition_Decision',FV['ACQ'],'Sheet 16 — Final acquisition decision and recommended re-acquisition conditions',widths={c:45 for c in FV['ACQ'].columns[1:]},wrap=True)
sheet('17_Final_Selection',FV['SEL'],'Sheet 17 — Final spectrum selection (reference groups). NAM: 15 (interim); LTA: 4 (< 10 → re-acquire); LPS: none',
 ['Selection rule: every spectrum acquired at the reference condition that passes QC is retained; none were dropped to improve statistics. The set is NOT filled with spectra from other conditions.'])
MMt=D['GRP'].copy(); MMt.columns=['Group','n','SNR mean','SNR SD','Pairwise r mean (glass zone excl.)','Pairwise r min','r(mean, median)','max |mean−median| / max(mean)','Fraction of points |mean−median| > 1 SD','Max |mean−median|/SD at markers','Median pointwise CV (%)']
sheet('18_Mean_vs_Median',MMt.round(4),'Sheet 18 — Mean vs median representative spectrum and group reproducibility',
 ['Decision rule: if r(mean, median) ≥ 0.999 and |mean − median| < 1 SD at every wavenumber and marker, the spectra are consistent and the MEAN is used (median reported alongside). Otherwise the median is used and outliers are investigated.',
  'NAM_REF and LTA_REF meet the rule → mean selected. LPS_ALL: r pairs down to −0.47 → no representative spectrum.'])
CT2=CT.copy(); CT2.columns=['Biomarker','Peak (nominal)','Position mean ± SD','Detected in reference','Min detection on any day (%) / LTA all-valid (%)','Detected in independent-spot set (%)','SNR','CV (%)','Prominence (counts)','Inside glass zone','Height retained after glass-template correction (%)','Distance from range edge (cm-1)','C1 reproducible','C2 SNR≥15 & CV≤15 %','C3 background-separated (not glass zone, ≥90 % retained, ≥75 cm-1 from edge)','C4 literature-plausible','Category']
sheet('19_Peak_Criteria',CT2.round(1),'Sheet 19 — Peak promotion criteria (primary = C1 + C2 + C3 + C4)',
 ['C1: detected in all reference spectra and ≥95 % of valid spectra on every day (NAM) or ≥8/9 valid spectra (LTA) and all independent spots (NAM). Thresholds are project conventions.','Supporting: detected in all reference spectra and ≥90 % overall, but fails at least one primary criterion. Uncertain: otherwise. Glass/background: inside glass zone and removed by glass correction.'])
sheet('20_Chemical_Assignment',FV['CHEM'],'Sheet 20 — Chemical / functional-group interpretation of final markers (TENTATIVE assignments)',
 ['NAM (MurNAc) = 2-acetamido-3-O-[(R)-1-carboxyethyl]-2-deoxy-D-glucopyranose: pyranose ring, N-acetyl amide, 3-O-lactyl ether, carboxylic acid. LTA = poly(glycerol-phosphate) chain with D-Ala/glycosyl substituents on a diacylglycerol glycolipid anchor.'],widths={'Molecular_vibration':38,'Chemical_component':32,'Biomarker_relevance':40},wrap=True)
BG=D['BG']; sheet('21_Background_Check',BG,'Sheet 21 — Background-origin screen for each final marker (glass, instrument, fluorescence, cosmic ray, contamination, processing)',widths={c:42 for c in BG.columns[2:]},wrap=True)
Q=D['QC'][['file','cls','power','int_s','acc','spot','snr','noise','max_ds','min_ds','max_raw_fp','max_raw_any','sat_laser','sat_fp','baseline_mean','baseline_range','baseline_slope_per100','dyn_range','prominence','spikes','r_group','qc','reason','in_ref']].copy()
Q.columns=['Spectrum','Class','Power (%)','Int (s)','Acc','Spot','SNR','Noise (counts, robust)','Max dark-sub. (400–1800)','Min dark-sub. (400–1800)','Max raw (400–1800)','Max raw (full)','Saturated laser line only','Saturated in 400–1800','Baseline mean','Baseline range','Baseline slope (counts/100 cm-1)','Dynamic range (max/noise)','Strongest-marker prominence','Spikes removed','Correlation with group mean','QC','Reason','In reference']
for c in Q.columns[6:21]:
  if Q[c].dtype.kind=='f': Q[c]=Q[c].round(2)
sheet('22_QC_Extended',Q,'Sheet 22 — Post-processing QC for every bulk/blank/suspect spectrum (pipeline v3)',
 ['Acceptance: not saturated in 400–1800 (16-bit detector limit, physical); SNR ≥ 10 (project convention); r to class mean ≥ 0.85 with glass zone excluded (project convention); no contamination pattern; no distortion (negative counts / abnormal dark). Spikes are corrected, not a rejection criterion.'],widths={'Reason':70})
S1=SE['Sm'].rename(columns={'setting':'Savitzky–Golay setting','height_vs_unsmoothed':'Marker height relative to unsmoothed','max_abs_shift':'Max peak-position shift (cm-1)'})
S2=SE['Al'].rename(columns={'setting':'ALS setting','mean':'Marker height / default (mean)','min':'min','max':'max'})
S3=SE['NR'].rename(columns={'window':'SG window (order 3)','noise_ratio':'Residual noise ratio (smoothed/unsmoothed)'})
ws=sheet('23_Preprocessing_Sensitivity',S1.round(3),'Sheet 23 — Sensitivity of marker heights to smoothing and baseline parameters (reference spectra)',
 ['SG 11/3 lowers the apex of the narrow NAM bands by ~16 % (positions within 1 pixel); LTA bands by ~3 %. ALS settings change NAM heights by < ±12 % but LTA heights by −40 % to +25 % (strong fluorescence). Absolute intensities are therefore pipeline-specific.'])
for df_,off in [(S2.round(3),len(S1)+7),(S3.round(3),len(S1)+len(S2)+11)]:
  for j,c in enumerate(df_.columns,1): x=ws.cell(off,j,c); x.font=H; x.fill=HF
  for i,row in enumerate(df_.itertuples(index=False),off+1):
    for j,v in enumerate(row,1): ws.cell(i,j,v.item() if isinstance(v,np.generic) else v).font=F
sheet('24_v3_Change_Log',FV['CHG'],'Sheet 24 — Pipeline v3 corrections and their effect on reported values',widths={'Reason':70,'v3':45,'v2.1':40},wrap=True)
TPL=pd.DataFrame([['NAM_Bulk_YYYYMMDD_P090pct_T30s_A05_S01.csv','YYYY-MM-DD','NAM','90','30','5','S01','slide ID / lot','spot position (x,y) or description','sample prep (e.g. powder pressed flat, amount)','objective (write it down)','operator','raw max 400–1800; saturation Y/N','notes'],
 ['Glass_Blank_YYYYMMDD_P090pct_T30s_A05_S01.csv','YYYY-MM-DD','Glass blank','90','30','5','S01','same slide type/lot','empty area','—','same objective','','','matched to NAM set']],
 columns=['File name','Date','Biomarker','Power (%)','Integration (s)','Accumulations','Spot/replicate','Glass slide info','Spot position','Sample preparation','Objective','Operator','Saturation check','Notes'])
sheet('25_Acquisition_Log_Template',TPL,'Sheet 25 — Acquisition log template for re-acquisition (fill one row per spectrum)')
# move central sheet next to README
wb.move_sheet('Traceability_Table',offset=1-wb.sheetnames.index('Traceability_Table'))
ws=wb['README']; r=ws.max_row+2
for t in ['VERSION 3 (23 Sep 2026): pipeline corrections (spike width guard, hot-pixel fix, robust noise) — see sheet 24. New sheets: Traceability_Table (central), 15–25.',
          'Decisions: NAM = GO for identity/markers + top-up re-acquisition (≥10 independent spots at 90 %/30 s/5 acc); LTA = RE-ACQUIRE (80 %/45 s/2 acc, 10–15 spots); LPS = RE-ACQUIRE (70/80/90 % screening, 30 s, 5 acc, matched glass blanks).']:
  ws.cell(r,1,t).font=Font(name='Arial',bold=True,size=10); r+=1
wb.save(fn); print(wb.sheetnames)
